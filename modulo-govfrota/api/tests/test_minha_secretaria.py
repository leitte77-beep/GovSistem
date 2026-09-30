"""Painel "Minha Secretaria": números só das secretarias do usuário."""

import json
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import pytest

from app.models.abastecimento import Abastecimento
from app.models.combustivel import Combustivel
from app.models.unidade import Unidade
from app.models.veiculo import Veiculo

API = "/api/govfrota"


async def _cenario(client, _db, make_tenant, setup_frota):
    admin = await make_tenant("ADMIN")
    org = admin["org"]
    f = await setup_frota(org)
    saude = Unidade(organization_id=org.id, nome="Saúde", sigla="SMS")
    educ = Unidade(organization_id=org.id, nome="Educação")
    _db.add_all([saude, educ])
    await _db.flush()
    carro = f["veiculo"]
    carro.unidade_id = saude.id
    onibus = Veiculo(organization_id=org.id, placa="EDU1A23", modelo="Onibus", tipo="ONIBUS",
                     quilometragem_atual=1000, situacao="DISPONIVEL", unidade_id=educ.id)
    arla = Combustivel(organization_id=org.id, nome="ARLA 32", categoria="FLUIDO_AUXILIAR")
    _db.add_all([onibus, arla])
    await _db.flush()
    agora = datetime.now(timezone.utc)

    def ab(v, un, litros, km, custo, comb=None, alertas=None, quando=agora):
        return Abastecimento(
            organization_id=org.id, veiculo_id=v.id, combustivel_id=(comb or f["combustivel"]).id,
            modalidade="TANQUE_PROPRIO", tanque_id=f["tanque"].id, unidade_id=un.id,
            quantidade_litros=Decimal(litros), quilometragem=km, data_abastecimento=quando,
            custo_total=Decimal(custo), origem="ADMIN", status="CONFIRMADO",
            alertas=json.dumps(alertas) if alertas else None,
        )

    _db.add_all([
        ab(carro, saude, "40", 50000, "240", quando=agora - timedelta(minutes=30)),
        ab(carro, saude, "50", 50500, "300", alertas=["SEM_CONTRATO", "LITROS_ACIMA_MEDIA"]),
        ab(carro, saude, "5", 50500, "20", comb=arla),
        ab(onibus, educ, "200", 1500, "1200"),
    ])
    await _db.commit()
    sec = await make_tenant("CONSULTA", org=org)
    r = await client.put(f"{API}/acessos/usuarios/{sec['user'].id}",
                         json={"perfil": "SECRETARIO", "unidade_ids": [str(saude.id)]}, headers=admin["headers"])
    assert r.status_code == 200, r.text
    return admin, sec, saude, educ, carro


@pytest.mark.asyncio
async def test_painel_do_secretario(client, _db, make_tenant, setup_frota):
    admin, sec, saude, educ, carro = await _cenario(client, _db, make_tenant, setup_frota)
    d = (await client.get(f"{API}/secretaria/painel", headers=sec["headers"])).json()

    assert [s["nome"] for s in d["secretarias"]] == ["Saúde"]
    i = d["indicadores"]
    assert i["gasto"] == 560.0 and i["litros"] == 95.0 and i["abastecimentos"] == 3
    assert i["veiculos_ativos"] == 1
    # 500 km com 90 L de combustível (ARLA fica fora do consumo).
    assert i["km_percorridos"] == 500 and i["consumo_km_l"] == 5.56
    assert i["custo_por_km"] == 1.12
    assert i["com_alerta"] == 1
    [v] = d["veiculos"]
    assert v["placa"] == carro.placa and v["secretaria"] == "Saúde" and v["km_percorridos"] == 500
    assert {c["combustivel"] for c in d["por_combustivel"]} == {"Diesel S10", "ARLA 32"}
    assert d["mensal"][-1]["gasto"] == 560.0 and len(d["mensal"]) == 12
    assert [a["codigo"] for a in d["alertas"][0]["alertas"]] == ["SEM_CONTRATO", "LITROS_ACIMA_MEDIA"]
    assert "caráter gerencial" in d["aviso"]

    # Pedir outra secretaria não fura o escopo.
    r = await client.get(f"{API}/secretaria/painel", params={"unidade_id": str(educ.id)}, headers=sec["headers"])
    assert r.status_code == 404


@pytest.mark.asyncio
async def test_gestor_escolhe_a_secretaria_e_periodo(client, _db, make_tenant, setup_frota):
    admin, sec, saude, educ, carro = await _cenario(client, _db, make_tenant, setup_frota)
    h = admin["headers"]
    geral = (await client.get(f"{API}/secretaria/painel", headers=h)).json()
    assert geral["indicadores"]["gasto"] == 1760.0 and not geral["restrito"]
    educacao = (await client.get(f"{API}/secretaria/painel", params={"unidade_id": str(educ.id)}, headers=h)).json()
    assert educacao["indicadores"]["litros"] == 200.0
    passado = (await client.get(f"{API}/secretaria/painel", params={"inicio": "2020-01-01", "fim": "2020-01-31"}, headers=h)).json()
    assert passado["indicadores"]["abastecimentos"] == 0
    r = await client.get(f"{API}/secretaria/painel", params={"inicio": "2020-02-01", "fim": "2020-01-01"}, headers=h)
    assert r.status_code == 422


@pytest.mark.asyncio
async def test_motoristas_da_secretaria_sem_dados_pessoais(client, _db, make_tenant, setup_frota):
    admin, sec, saude, educ, carro = await _cenario(client, _db, make_tenant, setup_frota)
    from sqlalchemy import update

    from app.models.motorista import Motorista

    mot = (await _db.execute(__import__("sqlalchemy").select(Motorista))).scalars().first()
    await _db.execute(update(Abastecimento).where(Abastecimento.veiculo_id == carro.id).values(motorista_id=mot.id))
    await _db.commit()
    d = (await client.get(f"{API}/secretaria/painel", headers=sec["headers"])).json()
    [m] = d["motoristas"]
    assert m["nome"] == mot.nome and m["abastecimentos"] == 3 and m["veiculos"] == [carro.placa]
    assert "cpf" not in m and "telefone" not in m
    # A tela de cadastro de motoristas continua fechada para o secretário.
    assert (await client.get(f"{API}/motoristas", headers=sec["headers"])).status_code == 403


@pytest.mark.asyncio
async def test_prestacao_de_contas_pdf_e_xlsx(client, _db, make_tenant, setup_frota):
    admin, sec, saude, educ, carro = await _cenario(client, _db, make_tenant, setup_frota)
    r = await client.get(f"{API}/secretaria/prestacao-contas", params={"formato": "pdf"}, headers=sec["headers"])
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    r = await client.get(f"{API}/secretaria/prestacao-contas", params={"formato": "xlsx"}, headers=sec["headers"])
    assert r.status_code == 200
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:3] == ["Indicadores", "Por combustível", "Veículos"]
    ind = wb["Indicadores"]
    assert ind["A6"].value == "Gasto no período (R$)" and ind["B6"].value == 560.0
    # Secretário não gera o da Educação.
    r = await client.get(f"{API}/secretaria/prestacao-contas", params={"unidade_id": str(educ.id)}, headers=sec["headers"])
    assert r.status_code == 404
    assert (await client.get(f"{API}/secretaria/prestacao-contas", params={"formato": "doc"}, headers=sec["headers"])).status_code == 422
