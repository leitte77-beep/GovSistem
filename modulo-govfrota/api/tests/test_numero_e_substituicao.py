"""Número do abastecimento (AB-ano/seq) e motivo obrigatório na troca de nota."""

import pytest

from tests.cenario_posto import _cenario
from tests.test_abastecimentos_redesenho import _abastecer_app, _login_app
from tests.test_portal_posto import API, P, _criar_acesso, _entrar

pytestmark = pytest.mark.asyncio


async def test_numero_sequencial_e_busca(client, make_tenant, setup_frota):
    tenant = await make_tenant("ADMIN")
    frota = await setup_frota(tenant["org"])
    app_h = await _login_app(client, frota)
    a1 = (await _abastecer_app(client, app_h, frota["veiculo"].id, 45, 50350)).json()
    a2 = (await _abastecer_app(client, app_h, frota["veiculo"].id, 30, 50600)).json()
    n1 = (await client.get(f"{API}/abastecimentos/{a1['id']}", headers=tenant["headers"])).json()["numero"]
    n2 = (await client.get(f"{API}/abastecimentos/{a2['id']}", headers=tenant["headers"])).json()["numero"]
    assert n1.startswith("AB-") and n1 != n2
    seq2 = int(n2.split("/")[1])
    assert seq2 == int(n1.split("/")[1]) + 1

    for termo in (n2, n2.split("-", 1)[1], str(seq2)):
        achados = (await client.get(f"{API}/abastecimentos", params={"search": termo}, headers=tenant["headers"])).json()
        assert a2["id"] in {a["id"] for a in achados}, termo
        if "/" in termo:
            assert a1["id"] not in {a["id"] for a in achados}, termo


async def test_substituir_nota_exige_motivo_e_fica_no_historico(client, _db, make_tenant, setup_frota):
    admin, _saude, _educ, posto, *_ = await _cenario(client, _db, make_tenant, setup_frota)
    acesso = await _criar_acesso(client, admin, posto, "posto.troca")
    h = await _entrar(client, "posto.troca", acesso["senha_provisoria"])
    alvo = (await client.get(f"{P}/abastecimentos", params={"nota": "PENDENTE"}, headers=h)).json()[0]
    assert alvo["numero"].startswith("AB-")
    url = f"{P}/abastecimentos/{alvo['id']}/nota"
    pdf = [("arquivos", ("d.pdf", b"%PDF-1 a", "application/pdf"))]

    assert (await client.post(url, files=pdf, headers=h)).status_code == 201
    # Segunda vez: sem motivo recusa; "OUTRO" sem descrição recusa.
    assert (await client.post(url, files=pdf, headers=h)).status_code == 422
    r = await client.post(url, files=pdf, data={"motivo": "OUTRO"}, headers=h)
    assert r.status_code == 422
    r = await client.post(url, files=pdf, data={"motivo": "ARQUIVO_ERRADO", "motivo_detalhe": "PDF cortado"}, headers=h)
    assert r.status_code == 201 and r.json()["substituida"] is True

    hist = (await client.get(f"{API}/notas-fiscais/abastecimento/{alvo['id']}/historico", headers=admin["headers"])).json()
    assert [(v["versao"], v["ativo"]) for v in hist] == [(1, False), (2, True)]
    assert hist[1]["motivo"] == "Arquivo errado ou ilegível" and hist[1]["motivo_detalhe"] == "PDF cortado"
    # A versão antiga continua acessível para consulta.
    assert (await client.get(f"{API}/notas-fiscais/{hist[0]['id']}/arquivo", headers=admin["headers"])).status_code == 200

    subst = (await client.get(f"{API}/notas-fiscais", params={"inicio": "2000-01-01", "situacao": "SUBSTITUIDA"}, headers=admin["headers"])).json()
    assert [i["id"] for i in subst["itens"]] == [alvo["id"]]


async def test_relatorio_mostra_nota_e_numero(client, _db, make_tenant, setup_frota):
    admin, _saude, _educ, posto, *_ = await _cenario(client, _db, make_tenant, setup_frota)
    acesso = await _criar_acesso(client, admin, posto, "posto.rel")
    h = await _entrar(client, "posto.rel", acesso["senha_provisoria"])
    alvo = (await client.get(f"{P}/abastecimentos", params={"nota": "PENDENTE"}, headers=h)).json()[0]
    pdf = [("arquivos", ("d.pdf", b"%PDF-1 a", "application/pdf"))]
    assert (await client.post(f"{P}/abastecimentos/{alvo['id']}/nota", files=pdf, headers=h)).status_code == 201

    url = f"{API}/relatorios/abastecimentos"
    rel = (await client.get(url, params={"data_inicio": "2000-01-01"}, headers=admin["headers"])).json()
    item = next(i for i in rel["itens"] if i["id"] == alvo["id"])
    # Só o PDF enviado: antes aparecia "—" na coluna NF.
    assert item["situacao_nota"] == "SO_PDF" and item["numero_nf"] == "PDF" and item["numero"].startswith("AB-")

    com = (await client.get(url, params={"data_inicio": "2000-01-01", "situacao_nota": "COM_NOTA"}, headers=admin["headers"])).json()
    assert [i["id"] for i in com["itens"]] == [alvo["id"]]
    sem = (await client.get(url, params={"data_inicio": "2000-01-01", "situacao_nota": "SEM_NOTA"}, headers=admin["headers"])).json()
    assert alvo["id"] not in {i["id"] for i in sem["itens"]}
    assert all(i["situacao_nota"] == "PENDENTE" for i in sem["itens"])

    for formato in ("csv", "xlsx", "pdf"):
        r = await client.get(url, params={"data_inicio": "2000-01-01", "formato": formato, "fornecedor_id": str(posto.id)}, headers=admin["headers"])
        assert r.status_code == 200, formato


async def test_relatorio_agrupado_com_subtotais(client, _db, make_tenant, setup_frota):
    admin, *_ = await _cenario(client, _db, make_tenant, setup_frota)
    url = f"{API}/relatorios/abastecimentos"
    base = {"data_inicio": "2000-01-01"}
    rel = (await client.get(url, params={**base, "agrupar_por": "secretaria"}, headers=admin["headers"])).json()
    grupos = rel["grupos"]
    assert len(grupos) >= 2
    assert sum(g["abastecimentos"] for g in grupos) == rel["total_registros"]
    assert round(sum(g["total"] for g in grupos), 2) == rel["total_gasto"]
    assert [g["total"] for g in grupos] == sorted((g["total"] for g in grupos), reverse=True)
    assert "itens" not in grupos[0]

    assert (await client.get(url, params={**base, "agrupar_por": "xyz"}, headers=admin["headers"])).status_code == 422
    for agrup in ("veiculo", "dia", "posto"):
        for formato in ("pdf", "xlsx"):
            r = await client.get(url, params={**base, "agrupar_por": agrup, "formato": formato}, headers=admin["headers"])
            assert r.status_code == 200, (agrup, formato)
    r = await client.get(url, params={**base, "agrupar_por": "secretaria", "formato": "xlsx"}, headers=admin["headers"])
    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Resumo"
