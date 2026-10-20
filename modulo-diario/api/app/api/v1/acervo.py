"""Acervo de atos: o livro de registro de tudo o que a entidade emitiu.

Não é uma tabela paralela — é uma visão organizada sobre as matérias, que já
guardam tipo, número, ano, data, entidade (unidade publicadora: Prefeitura ou
Câmara), conteúdo e situação. Assim um ato nunca existe em dois lugares.

Os atos se dividem em duas famílias:

* ``numerado``  — Portaria, Decreto, Lei, Resolução...: o número é do Diário,
  em sequência própria por entidade/tipo/ano.
* ``expediente`` — avisos de licitação, extratos, processos de compra,
  relatórios contábeis...: o número vem do processo/contrato/período.

A família vem de ``ActType.config["family"]`` quando configurada; senão, do
nome do tipo.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import uuid
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import Integer, and_, cast, extract, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.auth import get_current_user, require_roles
from app.core.database import get_db
from app.models.act_type import ActType
from app.models.edition import Edition
from app.models.edition_item import EditionItem
from app.models.enums import EditionStatus, MatterStatus
from app.models.matter import Matter
from app.models.org_unit import OrgUnit
from app.models.user import User
from app.services.edition_classification import entity_rank

router = APIRouter(prefix="/acervo", tags=["acervo"])

FAMILY_NUMBERED = "numerado"
FAMILY_EXPEDIENT = "expediente"
_NUMBERED_WORDS = {"portaria", "decreto", "lei", "resolucao", "deliberacao", "ato"}

# Situação "registrado": rascunho marcado como somente registro no acervo.
STATUS_REGISTERED = "registered"

STATUS_LABELS = {
    STATUS_REGISTERED: "Registrado",
    MatterStatus.DRAFT.value: "Rascunho",
    MatterStatus.REVIEW.value: "Em aprovação",
    MatterStatus.REJECTED.value: "Devolvido",
    MatterStatus.APPROVED.value: "Aguardando edição",
    MatterStatus.PUBLISHED.value: "Publicado",
    MatterStatus.ARCHIVED.value: "Arquivado",
}

_SORT_COLUMNS = {"numero", "data", "titulo", "atualizado"}
_EXPORT_LIMIT = 5000


def _normalize(value) -> str:
    text = str(value or "").strip().lower()
    return "".join(
        ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)
    )


def act_type_family(act_type: ActType) -> str:
    config = act_type.config if isinstance(act_type.config, dict) else {}
    explicit = str(config.get("family") or "").strip().lower()
    if explicit in (FAMILY_NUMBERED, FAMILY_EXPEDIENT):
        return explicit
    words = set(re.findall(r"[a-z]+", _normalize(act_type.name)))
    return FAMILY_NUMBERED if words & _NUMBERED_WORDS else FAMILY_EXPEDIENT


def _year_expr():
    """Ano do ato: o estruturado; nos legados, o da data do ato/publicação/criação."""
    return func.coalesce(
        Matter.act_year,
        cast(
            extract(
                "year",
                func.coalesce(
                    Matter.act_date,
                    func.date(Matter.published_at),
                    func.date(Matter.created_at),
                ),
            ),
            Integer,
        ),
    )


def _number_sort_exprs(desc: bool) -> list:
    """Ordem numérica sem cast (números legados podem ter letras): o mais
    curto vem antes ("04" < "235"), depois a ordem do texto."""
    length = func.length(Matter.act_number)
    if desc:
        return [length.desc().nulls_last(), Matter.act_number.desc().nulls_last()]
    return [length.asc().nulls_last(), Matter.act_number.asc().nulls_last()]


def _visibility(user: User):
    """Publicados são públicos: todos da organização veem. O resto segue a
    mesma regra da lista de matérias."""
    roles = {ur.role.name for ur in user.user_roles}
    if roles & {"ADMIN", "AUDITOR", "CONSULTA"}:
        return None
    own = Matter.author_id == user.id
    if "REVISOR" in roles:
        own = or_(own, Matter.status.in_([MatterStatus.REVIEW, MatterStatus.DRAFT]))
    if "DIAGRAMADOR" in roles:
        own = or_(own, Matter.status == MatterStatus.APPROVED)
    return or_(Matter.status == MatterStatus.PUBLISHED, own)


def _base_conditions(user: User) -> list:
    conditions = [Matter.organization_id == user.organization_id]
    visible = _visibility(user)
    if visible is not None:
        conditions.append(visible)
    return conditions


async def _family_type_ids(db: AsyncSession, family: str) -> list[uuid.UUID]:
    types = (await db.execute(select(ActType))).scalars().all()
    return [t.id for t in types if act_type_family(t) == family]


class AcervoFilters(BaseModel):
    act_type_id: Optional[uuid.UUID] = None
    org_unit_id: Optional[uuid.UUID] = None
    sem_entidade: bool = False
    family: Optional[str] = None
    year: Optional[int] = None
    status: Optional[str] = None
    q: Optional[str] = None
    date_from: Optional[date] = None
    date_to: Optional[date] = None


def _filters(
    act_type_id: Optional[uuid.UUID] = None,
    org_unit_id: Optional[uuid.UUID] = None,
    sem_entidade: bool = False,
    family: Optional[str] = Query(None, pattern="^(numerado|expediente)$"),
    year: Optional[int] = None,
    status: Optional[str] = None,
    q: Optional[str] = Query(None, max_length=200),
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
) -> AcervoFilters:
    if status and status not in STATUS_LABELS:
        raise HTTPException(422, "Situação inválida")
    return AcervoFilters(
        act_type_id=act_type_id, org_unit_id=org_unit_id, sem_entidade=sem_entidade,
        family=family, year=year, status=status, q=q, date_from=date_from, date_to=date_to,
    )


async def _conditions(db: AsyncSession, user: User, f: AcervoFilters) -> list:
    conditions = _base_conditions(user)
    if f.act_type_id:
        conditions.append(Matter.act_type_id == f.act_type_id)
    if f.sem_entidade:
        conditions.append(Matter.org_unit_id.is_(None))
    elif f.org_unit_id:
        conditions.append(Matter.org_unit_id == f.org_unit_id)
    if f.family:
        conditions.append(Matter.act_type_id.in_(await _family_type_ids(db, f.family)))
    if f.year:
        conditions.append(_year_expr() == f.year)
    if f.status == STATUS_REGISTERED:
        conditions += [Matter.status == MatterStatus.DRAFT, Matter.gazette_publication.is_(False)]
    elif f.status == MatterStatus.DRAFT.value:
        conditions += [Matter.status == MatterStatus.DRAFT, Matter.gazette_publication.is_not(False)]
    elif f.status:
        conditions.append(Matter.status == f.status)
    else:
        # Arquivado = descartado; só aparece quando pedido explicitamente.
        conditions.append(Matter.status != MatterStatus.ARCHIVED)
    act_date = func.coalesce(Matter.act_date, func.date(Matter.published_at), func.date(Matter.created_at))
    if f.date_from:
        conditions.append(act_date >= f.date_from)
    if f.date_to:
        conditions.append(act_date <= f.date_to)
    if f.q and f.q.strip():
        like = f"%{f.q.strip()}%"
        conditions.append(
            or_(
                Matter.title.ilike(like),
                Matter.summary.ilike(like),
                Matter.plain_text.ilike(like),
                Matter.act_number.ilike(like),
                Matter.responsible_name.ilike(like),
            )
        )
    return conditions


# ── Respostas ────────────────────────────────────────────────────────────────


class AcervoItem(BaseModel):
    id: uuid.UUID
    title: str
    summary: Optional[str]
    act_type_id: uuid.UUID
    act_type_name: str
    family: str
    org_unit_id: Optional[uuid.UUID]
    org_unit_name: Optional[str]
    act_number: Optional[str]
    act_year: Optional[int]
    act_date: Optional[date]
    responsible_name: Optional[str]
    status: str
    status_label: str
    published_at: Optional[datetime]
    edition_id: Optional[uuid.UUID] = None
    edition_number: Optional[int] = None
    edition_year: Optional[int] = None
    edition_date: Optional[date] = None
    updated_at: datetime


class AcervoPage(BaseModel):
    items: list[AcervoItem]
    total: int


class AcervoFolder(BaseModel):
    act_type_id: uuid.UUID
    name: str
    family: str
    total: int


class AcervoEntity(BaseModel):
    id: Optional[uuid.UUID]
    name: str
    total: int


class AcervoFolders(BaseModel):
    folders: list[AcervoFolder]
    entities: list[AcervoEntity]
    years: list[int]
    total: int


async def _rows(
    db: AsyncSession,
    conditions: list,
    *,
    sort: str = "numero",
    order: str = "desc",
    skip: int = 0,
    limit: int = 50,
) -> list[AcervoItem]:
    year = _year_expr().label("ano")
    query = (
        select(Matter, ActType, OrgUnit, year)
        .join(ActType, ActType.id == Matter.act_type_id)
        .outerjoin(OrgUnit, OrgUnit.id == Matter.org_unit_id)
        .where(and_(*conditions))
    )
    desc = order.lower() != "asc"

    def ordered(col):
        return col.desc().nulls_last() if desc else col.asc().nulls_last()

    act_date = func.coalesce(Matter.act_date, func.date(Matter.published_at), func.date(Matter.created_at))
    if sort == "titulo":
        query = query.order_by(ordered(Matter.title))
    elif sort == "data":
        # No mesmo dia, o número; atos sem número (expediente sem nº de
        # processo) pela ordem de cadastro — "mais recentes" mostra primeiro
        # o que acabou de entrar.
        query = query.order_by(ordered(act_date), *_number_sort_exprs(desc), ordered(Matter.created_at))
    elif sort == "atualizado":
        query = query.order_by(ordered(Matter.updated_at))
    else:
        query = query.order_by(ordered(year), *_number_sort_exprs(desc), ordered(act_date))
    query = query.order_by(Matter.id).offset(skip).limit(limit)

    rows = (await db.execute(query)).all()

    # Primeira edição publicada de cada ato da página (consulta à parte para
    # não multiplicar linhas quando o ato saiu em mais de uma edição).
    editions: dict[uuid.UUID, Edition] = {}
    matter_ids = [row[0].id for row in rows]
    if matter_ids:
        published = await db.execute(
            select(EditionItem.matter_id, Edition)
            .join(Edition, Edition.id == EditionItem.edition_id)
            .where(
                EditionItem.matter_id.in_(matter_ids),
                Edition.status == EditionStatus.PUBLISHED,
            )
            .order_by(Edition.publication_date.desc())
        )
        for matter_id, edition in published.all():
            editions[matter_id] = edition  # a última atribuição é a mais antiga

    items = []
    for row in rows:
        matter, act_type, org_unit, ano = row[0], row[1], row[2], row[3]
        edition = editions.get(matter.id)
        status = matter.status.value if hasattr(matter.status, "value") else str(matter.status)
        if status == MatterStatus.DRAFT.value and matter.gazette_publication is False:
            status = STATUS_REGISTERED
        items.append(
            AcervoItem(
                id=matter.id,
                title=matter.title,
                summary=matter.summary,
                act_type_id=act_type.id,
                act_type_name=act_type.name,
                family=act_type_family(act_type),
                org_unit_id=org_unit.id if org_unit else None,
                org_unit_name=org_unit.name if org_unit else None,
                act_number=matter.act_number,
                act_year=matter.act_year or ano,
                act_date=matter.act_date,
                responsible_name=matter.responsible_name,
                status=status,
                status_label=STATUS_LABELS.get(status, status),
                published_at=matter.published_at,
                edition_id=edition.id if edition else None,
                edition_number=edition.number if edition else None,
                edition_year=edition.year if edition else None,
                edition_date=edition.publication_date if edition else None,
                updated_at=matter.updated_at,
            )
        )
    return items


# ── Rotas ────────────────────────────────────────────────────────────────────


@router.get("/pastas", response_model=AcervoFolders)
async def acervo_folders(
    org_unit_id: Optional[uuid.UUID] = None,
    year: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Pastas por tipo de ato com contagem, entidades e anos disponíveis."""
    base = _base_conditions(user) + [Matter.status != MatterStatus.ARCHIVED]

    scoped = list(base)
    if org_unit_id:
        scoped.append(Matter.org_unit_id == org_unit_id)
    if year:
        scoped.append(_year_expr() == year)

    counts = dict(
        (
            await db.execute(
                select(Matter.act_type_id, func.count()).where(and_(*scoped)).group_by(Matter.act_type_id)
            )
        ).all()
    )
    types = (
        await db.execute(
            select(ActType).where(ActType.deleted_at.is_(None)).order_by(ActType.name)
        )
    ).scalars().all()
    folders = [
        AcervoFolder(act_type_id=t.id, name=t.name, family=act_type_family(t), total=counts.get(t.id, 0))
        for t in types
        if t.is_active or counts.get(t.id)
    ]

    year_scoped = list(base) + ([_year_expr() == year] if year else [])
    entity_counts = dict(
        (
            await db.execute(
                select(Matter.org_unit_id, func.count()).where(and_(*year_scoped)).group_by(Matter.org_unit_id)
            )
        ).all()
    )
    units = (
        await db.execute(
            select(OrgUnit).where(
                OrgUnit.organization_id == user.organization_id,
                OrgUnit.deleted_at.is_(None),
            )
        )
    ).scalars().all()
    units = sorted(units, key=lambda u: (entity_rank(u), _normalize(u.name)))
    entities = [AcervoEntity(id=u.id, name=u.name, total=entity_counts.get(u.id, 0)) for u in units]
    if entity_counts.get(None):
        entities.append(AcervoEntity(id=None, name="Sem entidade", total=entity_counts[None]))

    years = sorted(
        {
            int(y)
            for (y,) in (
                await db.execute(select(_year_expr()).where(and_(*base)).distinct())
            ).all()
            if y
        },
        reverse=True,
    )
    return AcervoFolders(
        folders=folders,
        entities=entities,
        years=years,
        total=sum(counts.values()),
    )


@router.get("", response_model=AcervoPage)
async def list_acervo(
    filters: AcervoFilters = Depends(_filters),
    sort: str = "numero",
    order: str = "desc",
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if sort not in _SORT_COLUMNS:
        raise HTTPException(422, "Ordenação inválida")
    conditions = await _conditions(db, user, filters)
    total = await db.scalar(select(func.count()).select_from(Matter).where(and_(*conditions)))
    items = await _rows(db, conditions, sort=sort, order=order, skip=skip, limit=limit)
    return AcervoPage(items=items, total=int(total or 0))


class AcervoSuggestion(BaseModel):
    family: str
    act_number: Optional[str]
    act_year: int
    act_date: date
    title: str
    responsible_id: Optional[uuid.UUID] = None
    responsible_name: Optional[str] = None
    responsible_role: Optional[str] = None
    responsible_source: Optional[str] = None  # "autoridade" | "ultimo_ato"


def _suggested_title(act_type: ActType, number: Optional[str], year: int, act_date: date) -> str:
    from app.core.act_type_config import format_act_title

    config = act_type.config if isinstance(act_type.config, dict) else {}
    if config.get("title_pattern") and number:
        title = format_act_title(
            config.get("title_pattern"), type_name=act_type.name, number=number, year=year, act_date=act_date
        )
        if title:
            return title
    if number:
        return f"{act_type.name.upper()} Nº {number}/{year}"
    return act_type.name.upper()


@router.get("/sugestoes", response_model=AcervoSuggestion)
async def acervo_suggestions(
    act_type_id: uuid.UUID,
    org_unit_id: Optional[uuid.UUID] = None,
    year: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Tudo o que o cadastro de um ato novo pode preencher sozinho.

    * número: próximo da série (tipo + entidade + ano) para atos numerados; nos
      de expediente o número é do processo/contrato e fica para o usuário;
    * responsável: a autoridade ativa da entidade; sem cadastro, o último
      responsável usado naquele tipo e entidade.
    """
    from app.models.authority import Authority
    from app.services.document_numbering import highest_used_number, institutional_today

    act_type = await db.get(ActType, act_type_id)
    if act_type is None or act_type.deleted_at is not None:
        raise HTTPException(404, "Tipo de ato não encontrado")
    if org_unit_id is not None:
        unit = await db.get(OrgUnit, org_unit_id)
        if unit is None or unit.organization_id != user.organization_id:
            raise HTTPException(404, "Entidade não encontrada")

    today = institutional_today()
    year = year or today.year
    family = act_type_family(act_type)
    number = None
    if family == FAMILY_NUMBERED:
        highest = await highest_used_number(
            db, organization_id=user.organization_id, act_type_id=act_type_id, year=year, org_unit_id=org_unit_id
        )
        number = f"{highest + 1:02d}"

    out = AcervoSuggestion(
        family=family,
        act_number=number,
        act_year=year,
        act_date=today,
        title=_suggested_title(act_type, number, year, today),
    )

    authorities = (
        await db.execute(
            select(Authority)
            .where(
                Authority.organization_id == user.organization_id,
                Authority.is_active.is_(True),
                Authority.org_unit_id.is_not_distinct_from(org_unit_id),
            )
            .order_by(Authority.created_at)
        )
    ).scalars().all()
    current = [a for a in authorities if not a.valid_until or a.valid_until >= today]
    if current:
        chosen = current[0]
        out.responsible_id, out.responsible_name, out.responsible_role = chosen.id, chosen.name, chosen.role
        out.responsible_source = "autoridade"
        return out

    last = (
        await db.execute(
            select(Matter.responsible_name, Matter.responsible_role)
            .where(
                Matter.organization_id == user.organization_id,
                Matter.act_type_id == act_type_id,
                Matter.org_unit_id.is_not_distinct_from(org_unit_id),
                Matter.responsible_name.is_not(None),
                Matter.status != MatterStatus.ARCHIVED,
            )
            .order_by(Matter.created_at.desc())
            .limit(1)
        )
    ).first()
    if last:
        out.responsible_name, out.responsible_role = last
        out.responsible_source = "ultimo_ato"
    return out


# ── Exportação ───────────────────────────────────────────────────────────────

_EXPORT_HEADERS = [
    "Entidade", "Tipo", "Número", "Ano", "Data", "Título", "Súmula",
    "Responsável", "Situação", "Edição", "Data da publicação",
]


def _display_date(item: AcervoItem) -> str:
    """Data do ato; nos legados sem data própria, a da publicação."""
    if item.act_date:
        return item.act_date.strftime("%d/%m/%Y")
    if item.edition_date:
        return item.edition_date.strftime("%d/%m/%Y") + " (publ.)"
    return ""


def _export_row(item: AcervoItem) -> list:
    edition = f"{item.edition_number}/{item.edition_year}" if item.edition_number else ""
    return [
        item.org_unit_name or "",
        item.act_type_name,
        item.act_number or "",
        item.act_year or "",
        _display_date(item),
        item.title,
        item.summary or "",
        item.responsible_name or "",
        item.status_label,
        edition,
        item.edition_date.strftime("%d/%m/%Y") if item.edition_date else "",
    ]


def _export_title(filters: AcervoFilters, items: list[AcervoItem]) -> str:
    parts = ["Acervo de atos"]
    if filters.act_type_id and items:
        parts = [items[0].act_type_name]
    if filters.org_unit_id and items:
        parts.append(items[0].org_unit_name or "")
    if filters.year:
        parts.append(str(filters.year))
    return " — ".join(p for p in parts if p)


def _xlsx(title: str, items: list[AcervoItem]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Atos"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} — {len(items)} ato(s)"])
    ws.append([])
    ws.append(_EXPORT_HEADERS)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for item in items:
        ws.append(_export_row(item))
    widths = [22, 20, 10, 8, 13, 45, 60, 28, 18, 12, 16]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(_EXPORT_HEADERS))}{ws.max_row}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _csv(items: list[AcervoItem]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow(_EXPORT_HEADERS)
    for item in items:
        writer.writerow(_export_row(item))
    # BOM para o Excel abrir acentuação corretamente.
    return ("﻿" + out.getvalue()).encode("utf-8")


def _pdf(title: str, organization_name: str, items: list[AcervoItem]) -> bytes:
    from html import escape

    from weasyprint import HTML

    rows = "".join(
        "<tr>" + "".join(f"<td>{escape(str(v))}</td>" for v in (
            r[0], r[1], f"{r[2]}/{r[3]}" if r[2] else "", r[4], r[5], r[8], r[9],
        )) + "</tr>"
        for r in (_export_row(item) for item in items)
    )
    html = f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: A4 landscape; margin: 14mm 12mm;
  @bottom-right {{ content: "Página " counter(page) " de " counter(pages); font-size: 8pt; color: #555; }} }}
body {{ font-family: "DejaVu Sans", sans-serif; font-size: 8.5pt; color: #111; }}
h1 {{ font-size: 13pt; margin: 0 0 2px; }}
.meta {{ color: #555; margin-bottom: 10px; }}
table {{ width: 100%; border-collapse: collapse; }}
thead {{ display: table-header-group; }}
th {{ background: #1f3a5f; color: #fff; text-align: left; padding: 4px; }}
td {{ border-bottom: 0.5pt solid #ccc; padding: 4px; vertical-align: top; }}
tr {{ page-break-inside: avoid; }}
</style></head><body>
<h1>{escape(title)}</h1>
<div class="meta">{escape(organization_name)} · gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} · {len(items)} ato(s)</div>
<table><thead><tr><th>Entidade</th><th>Tipo</th><th>Nº</th><th>Data</th><th style="width:40%">Título</th><th>Situação</th><th>Edição</th></tr></thead>
<tbody>{rows}</tbody></table></body></html>"""
    return HTML(string=html).write_pdf()


@router.get("/exportar")
async def export_acervo(
    formato: str = Query("xlsx", pattern="^(xlsx|csv|pdf)$"),
    filters: AcervoFilters = Depends(_filters),
    sort: str = "numero",
    order: str = "asc",
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if sort not in _SORT_COLUMNS:
        raise HTTPException(422, "Ordenação inválida")
    conditions = await _conditions(db, user, filters)
    total = await db.scalar(select(func.count()).select_from(Matter).where(and_(*conditions)))
    if (total or 0) > _EXPORT_LIMIT:
        raise HTTPException(
            422, f"A exportação está limitada a {_EXPORT_LIMIT} atos; refine os filtros."
        )
    items = await _rows(db, conditions, sort=sort, order=order, limit=_EXPORT_LIMIT)
    if sort == "numero":
        # Um livro de registro: Prefeitura antes da Câmara, cada tipo junto, e
        # dentro do tipo a sequência numérica (a ordem da consulta já é por número).
        entity_order = {
            u.id: entity_rank(u)
            for u in (
                await db.execute(select(OrgUnit).where(OrgUnit.organization_id == user.organization_id))
            ).scalars()
        }
        items.sort(key=lambda i: (entity_order.get(i.org_unit_id, 2), _normalize(i.act_type_name)))
    title = _export_title(filters, items)
    stamp = date.today().strftime("%Y%m%d")
    if formato == "csv":
        content, media = _csv(items), "text/csv; charset=utf-8"
    elif formato == "pdf":
        from app.models.organization import Organization

        organization = await db.get(Organization, user.organization_id)
        content, media = _pdf(title, getattr(organization, "name", "") or "", items), "application/pdf"
    else:
        content = _xlsx(title, items)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(
        content=content,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="acervo-atos-{stamp}.{formato}"'},
    )


# ── Visualizar / baixar ──────────────────────────────────────────────────────


async def _visible_matter(db: AsyncSession, user: User, matter_id: uuid.UUID) -> Matter:
    matter = (
        await db.execute(select(Matter).where(Matter.id == matter_id, *_base_conditions(user)))
    ).scalar_one_or_none()
    if matter is None:
        raise HTTPException(404, "Ato não encontrado")
    return matter


async def _published_edition(db: AsyncSession, matter_id: uuid.UUID) -> Optional[Edition]:
    """Primeira edição publicada em que o ato saiu (com assinaturas carregadas)."""
    from sqlalchemy.orm import selectinload

    return (
        await db.execute(
            select(Edition)
            .join(EditionItem, EditionItem.edition_id == Edition.id)
            .where(EditionItem.matter_id == matter_id, Edition.status == EditionStatus.PUBLISHED)
            .options(selectinload(Edition.signatures))
            .order_by(Edition.publication_date)
            .limit(1)
        )
    ).scalar_one_or_none()


async def _frozen_item(db: AsyncSession, edition: Edition, matter_id: uuid.UUID) -> Optional[dict]:
    """O ato exatamente como foi publicado (snapshot imutável da edição)."""
    from app.models.edition_publication_snapshot import EditionPublicationSnapshot

    snapshot = (
        await db.execute(
            select(EditionPublicationSnapshot)
            .where(EditionPublicationSnapshot.edition_id == edition.id, EditionPublicationSnapshot.is_valid.is_(True))
            .order_by(EditionPublicationSnapshot.frozen_at.desc())
            .limit(1)
        )
    ).scalar_one_or_none()
    if snapshot is None:
        return None
    for item in (snapshot.content or {}).get("items", []):
        if str(item.get("id")) == str(matter_id):
            return item
    return None


def _act_pdf(item: dict, *, organization, entity: str, act_type: str, edition: Optional[Edition]) -> bytes:
    from html import escape

    from weasyprint import HTML

    from app.services.edition_pdf import _localize_matter_images, _render_semantic_content
    from app.services.print_tables import prepare_tables_for_print
    from app.services.publication_title import display_title

    content = _render_semantic_content(item)
    if content is None:
        content = item.get("content_html") or ""
    content = _localize_matter_images(prepare_tables_for_print(content))

    responsible = item.get("responsible") or {}
    if edition is not None:
        footer = (
            f"Publicado no Diário Oficial — Edição nº {edition.number}/{edition.year}, "
            f"de {edition.publication_date.strftime('%d/%m/%Y')}. "
            "O documento oficial é a edição assinada digitalmente."
        )
        watermark = ""
    else:
        footer = "Documento não publicado — cópia para consulta."
        watermark = '<div class="wm">NÃO PUBLICADO</div>'
    summary = item.get("summary") if not item.get("semantic") else None
    signature = ""
    if responsible.get("name"):
        signature = (
            f'<div class="sig"><strong>{escape(responsible["name"])}</strong>'
            f'<br>{escape(responsible.get("role") or "")}</div>'
        )
    html = f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 22mm 20mm 20mm;
  @bottom-center {{ content: "{escape(footer)}"; font-size: 7.5pt; color: #555; }}
  @bottom-right {{ content: counter(page) "/" counter(pages); font-size: 7.5pt; color: #555; }} }}
body {{ font-family: "DejaVu Serif", serif; font-size: 11pt; line-height: 1.45; color: #111; }}
header {{ border-bottom: 1pt solid #1f3a5f; padding-bottom: 6px; margin-bottom: 16px; }}
header .org {{ font-family: "DejaVu Sans", sans-serif; font-weight: bold; font-size: 11pt; color: #1f3a5f; }}
header .meta {{ font-family: "DejaVu Sans", sans-serif; font-size: 8.5pt; color: #555; }}
h1 {{ font-size: 12.5pt; text-align: center; margin: 0 0 12px; }}
.summary {{ margin: 0 0 14px 40%; font-size: 10pt; text-align: justify; }}
.content p {{ text-align: justify; margin: 0 0 8px; }}
.content img {{ max-width: 100%; }}
.content table {{ border-collapse: collapse; width: 100%; }}
.content td, .content th {{ border: 0.5pt solid #444; padding: 3px; font-size: 9.5pt; }}
.sig {{ margin-top: 36px; text-align: center; }}
.wm {{ position: fixed; top: 40%; left: 0; right: 0; text-align: center; font-size: 54pt; color: rgba(200,0,0,0.10); transform: rotate(-30deg); font-family: "DejaVu Sans", sans-serif; }}
</style></head><body>{watermark}
<header><div class="org">{escape(getattr(organization, "name", "") or "")}</div>
<div class="meta">{escape(entity)} · {escape(act_type)}</div></header>
<h1>{escape(display_title(item))}</h1>
{f'<div class="summary">{escape(summary)}</div>' if summary else ''}
<div class="content">{content}</div>{signature}
</body></html>"""
    return HTML(string=html, base_url="/").write_pdf()


def _filename(matter: Matter, act_type: ActType, suffix: str = "") -> str:
    base = _normalize(act_type.name).replace(" ", "-")
    number = re.sub(r"[^0-9A-Za-z]+", "", matter.act_number or "") or "sn"
    year = matter.act_year or ""
    return f"{base}-{number}-{year}{suffix}.pdf".replace("--", "-")


@router.get("/{matter_id}/pdf")
async def acervo_act_pdf(
    matter_id: uuid.UUID,
    download: bool = False,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """PDF do ato sozinho. Publicado: o texto exatamente como saiu na edição."""
    from app.models.organization import Organization
    from app.semantic.snapshot import matter_snapshot

    matter = await _visible_matter(db, user, matter_id)
    act_type = await db.get(ActType, matter.act_type_id)
    unit = await db.get(OrgUnit, matter.org_unit_id) if matter.org_unit_id else None
    organization = await db.get(Organization, user.organization_id)
    edition = await _published_edition(db, matter.id)

    item = await _frozen_item(db, edition, matter.id) if edition else None
    if item is None:
        item = matter_snapshot(matter, position=0, section_title=None, attachments_override=[])

    content = _act_pdf(
        item,
        organization=organization,
        entity=unit.name if unit else "",
        act_type=act_type.name if act_type else "",
        edition=edition,
    )
    disposition = "attachment" if download else "inline"
    return Response(
        content=content,
        media_type="application/pdf",
        headers={"Content-Disposition": f'{disposition}; filename="{_filename(matter, act_type)}"'},
    )


@router.get("/{matter_id}/edicao")
async def acervo_edition_pdf(
    matter_id: uuid.UUID,
    download: bool = False,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """PDF oficial (assinado, quando houver) da edição em que o ato foi publicado."""
    from app.api.v1.public import _public_pdf_path, _resolve_upload_path
    from app.models.organization import Organization

    await _visible_matter(db, user, matter_id)
    edition = await _published_edition(db, matter_id)
    if edition is None:
        raise HTTPException(404, "Este ato ainda não foi publicado em edição.")
    organization = await db.get(Organization, user.organization_id)
    slug = getattr(organization, "slug", None)
    relative = _public_pdf_path(edition, slug)
    path = _resolve_upload_path(relative, slug) if relative else None
    if path is None:
        raise HTTPException(404, "O arquivo da edição não foi encontrado.")
    disposition = "attachment" if download else "inline"
    filename = f"diario-oficial-{edition.year}-{edition.number:04d}.pdf"
    return Response(
        content=path.read_bytes(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'{disposition}; filename="{filename}"'},
    )


# ── Importar Word ────────────────────────────────────────────────────────────


class WordImport(BaseModel):
    content_html: str
    text_boxes: int
    tables: int
    images_skipped: int
    warnings: list[str]


@router.post("/importar-word", response_model=WordImport)
async def acervo_import_word(
    file: UploadFile,
    user: User = Depends(require_roles("AUTOR", "ADMIN")),
):
    """Lê um .docx e devolve o HTML para o editor do Novo ato (nada é gravado).

    Diferente de colar: o conteúdo de caixas de texto/formas (quadros de
    vencedores em avisos de licitação) vem como texto e tabela editáveis.
    """
    from app.core.file_validator import validate_upload
    from app.services.docx_html import docx_to_html

    ext, content = await validate_upload(file)
    if ext != ".docx":
        raise HTTPException(422, "Envie um arquivo do Word no formato .docx.")
    try:
        from app.providers.antivirus import get_virus_scanner

        scan = await get_virus_scanner().scan(content, file.filename or "documento.docx")
        if not scan.clean:
            raise HTTPException(422, f"Arquivo recusado pelo antivírus: {scan.message}")
        result = docx_to_html(content)
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001 — arquivo corrompido/protegido
        raise HTTPException(422, "Não foi possível ler este arquivo do Word.") from exc
    return WordImport(
        content_html=result.html,
        text_boxes=result.text_boxes,
        tables=result.tables,
        images_skipped=result.images_skipped,
        warnings=result.warnings,
    )
