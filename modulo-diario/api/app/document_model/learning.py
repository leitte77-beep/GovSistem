"""Aprender com documentos existentes (Fase 4).

A IA analisa documentos oficiais reais (PDF/DOCX) e **propõe** um
``DocumentModelConfig`` (campos + seções), classificando textos fixos, campos
variáveis, blocos condicionais e assinatura. A proposta nunca é aplicada
automaticamente: o administrador revisa e aprova; só então uma nova versão
(rascunho) é criada pelo fluxo normal de versionamento.

A IA é assistente, não autora: o modelo ativo continua sendo a fonte de verdade.
"""

from __future__ import annotations

import io
import logging
import re
import uuid
from collections import Counter

from pydantic import ValidationError

from app.document_model.schemas import (
    SCOPE_DOCUMENT_TYPES,
    DocumentModelConfig,
    FieldType,
    SectionKind,
    extract_markers,
)
from app.services.ai.deepseek_client import DeepSeekClient
from app.services.ai.errors import AiInvalidResponseError

logger = logging.getLogger(__name__)

LEARN_PROMPT_VERSION = "dm-learn-v3"

# Limite de caracteres por documento e total, para caber no orçamento da IA.
MAX_CHARS_PER_DOC = 12000
MAX_TOTAL_CHARS = 40000


def _representative_excerpt(text: str, limit: int = MAX_CHARS_PER_DOC) -> str:
    """Preserva a cobertura do documento inteiro, não apenas a primeira página.

    Edições do Diário Oficial reúnem vários atos; cortar somente o início fazia
    a IA ignorar sistematicamente matérias do final, anexos e relatórios.
    """
    normalized = text.strip()
    if len(normalized) <= limit:
        return normalized
    head = limit // 3
    middle = limit // 3
    tail = limit - head - middle
    start = normalized[:head].rstrip()
    center_start = max(0, (len(normalized) - middle) // 2)
    center = normalized[center_start : center_start + middle].strip()
    end = normalized[-tail:].lstrip()
    return (
        f"{start}\n\n[... trecho intermediário preservado ...]\n\n"
        f"{center}\n\n[... trecho final preservado ...]\n\n{end}"
    )

_MARKER_KEY = re.compile(r"^[a-z0-9_]+$")

_KINDS = ", ".join(sorted(k.value for k in SectionKind))
_FIELD_TYPES = ", ".join(sorted(t.value for t in FieldType))


def extract_text(filename: str, content: bytes) -> str:
    """Extrai texto de um documento de referência (TXT, PDF ou DOCX)."""
    lower = (filename or "").lower()
    if lower.endswith(".txt"):
        try:
            return content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Arquivo TXT deve estar codificado em UTF-8.") from exc
    if lower.endswith(".pdf"):
        from pdfminer.high_level import extract_text as pdf_extract_text

        return pdf_extract_text(io.BytesIO(content)) or ""
    if lower.endswith(".docx"):
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError("Formato não suportado para treinamento (use PDF ou DOCX).")


def extract_visual_profile(filename: str, content: bytes) -> str:
    """Resume sinais visuais disponíveis no arquivo, sem enviar binário à IA."""
    lower = (filename or "").lower()
    if lower.endswith(".pdf"):
        from pdfminer.high_level import extract_pages
        from pdfminer.layout import LTChar, LTTextLine

        fonts: Counter[str] = Counter()
        sizes: Counter[float] = Counter()
        alignments: Counter[str] = Counter()
        upper = bold = italic = 0
        lines = 0
        def iter_lines(node):
            if isinstance(node, LTTextLine):
                yield node
            elif hasattr(node, "__iter__"):
                for child in node:
                    yield from iter_lines(child)

        for page in extract_pages(io.BytesIO(content)):
            page_width = float(getattr(page, "width", 595) or 595)
            for element in iter_lines(page):
                text = element.get_text().strip()
                if not text:
                    continue
                lines += 1
                x0, x1 = float(element.x0), float(element.x1)
                center = (x0 + x1) / 2
                if abs(center - page_width / 2) < page_width * 0.08:
                    alignments["center"] += 1
                elif x0 < page_width * 0.18:
                    alignments["left"] += 1
                else:
                    alignments["other_or_justified"] += 1
                chars = [obj for obj in element if isinstance(obj, LTChar)]
                for char in chars:
                    font = str(char.fontname or "unknown")
                    fonts[font] += 1
                    sizes[round(float(char.size), 1)] += 1
                    name = font.lower()
                    bold += int("bold" in name or "bd" in name)
                    italic += int("italic" in name or "oblique" in name)
                letters = [c for c in text if c.isalpha()]
                upper += int(bool(letters) and sum(c.isupper() for c in letters) / len(letters) >= 0.8)
        if not lines:
            return ""
        return (
            "\n[PERFIL VISUAL EXTRAÍDO DO PDF — use como evidência, não invente estilos]\n"
            f"fontes predominantes: {fonts.most_common(5)}; tamanhos: {sizes.most_common(5)}; "
            f"alinhamentos observados: {alignments.most_common()}; linhas em caixa alta: {upper}/{lines}; "
            f"caracteres em fontes com indício de negrito: {bold}; itálico: {italic}.\n"
        )
    if lower.endswith(".docx"):
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(content))
        fonts: Counter[str] = Counter()
        sizes: Counter[str] = Counter()
        bold = italic = 0
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if run.font.name:
                    fonts[run.font.name] += 1
                if run.font.size:
                    sizes[str(round(run.font.size.pt, 1))] += 1
                bold += int(run.bold is True)
                italic += int(run.italic is True)
        return (
            "\n[PERFIL VISUAL EXTRAÍDO DO DOCX — use como evidência]\n"
            f"fontes: {fonts.most_common(5)}; tamanhos: {sizes.most_common(5)}; "
            f"runs em negrito: {bold}; runs em itálico: {italic}.\n"
        )
    return ""


def build_system_message() -> str:
    return (
        "Você analisa documentos oficiais de uma prefeitura e propõe um MODELO "
        "documental estruturado. Responda SOMENTE com um objeto JSON, sem "
        "comentários, no formato:\n"
        '{"purpose": str, "scope_document_type": str, "document_title": str, '
        '"summary": str, "fields": [{"key": str, "label": str, "type": str, '
        '"required": bool, "options": [str]}], "sections": [{"id": str, '
        '"kind": str, "text": str, "alignment": str, "level": int, '
        '"number": str|null, "fixed_text": bool, "locked": bool, '
        '"font_weight": "normal|bold", "italic": bool, "uppercase": bool, '
        '"indent_em": number, "space_before_mm": number, "space_after_mm": number, '
        '"ai_generated": bool, "when_field": str|null, "when_value": str|null, '
        '"table_headers": [str], "table_rows": [[str]], '
        '"table_column_widths": [number], '
        '"entries": [{"name": str, "role": str}]}]}\n'
        f"Kinds permitidos: {_KINDS}.\n"
        f"Tipos de campo permitidos: {_FIELD_TYPES}.\n"
        "Regras obrigatórias:\n"
        "- Todo marcador {{chave}} no texto deve ter um campo com essa key.\n"
        "- Use {{chave}} para dados variáveis (nomes, datas, números, cargos).\n"
        "- Marque textos institucionais/legais invariáveis com \"fixed_text\": true "
        'e "locked": true (ex.: "RESOLVE:", "Registre-se e Publique-se").\n'
        "- Use \"when_field\"/\"when_value\" para trechos condicionais.\n"
        "- Prefira kind \"article\" para dispositivos, \"signature_block\" para "
        "assinatura, \"command\" para RESOLVE/DECRETA.\n"
        "- Preserve a forma específica de cada tipo de ato, sem aplicar um modelo "
        "genérico: leis e decretos devem manter capítulos, artigos, parágrafos, "
        "incisos e alíneas; portarias e resoluções devem manter a fórmula de "
        "comando (RESOLVE/RESOLVEM) e a numeração observada; editais e licitações "
        "devem manter identificação do processo, objeto, modalidade, datas, "
        "seções, subitens, tabelas e anexos; contratos, termos e extratos devem "
        "manter partes, vigência, valores, dotações, assinaturas e referências; "
        "relatórios, laudos e audiências devem manter títulos, quesitos, tabelas, "
        "conclusões e convocação.\n"
        "- Conserve exatamente a convenção de numeração observada em cada fonte "
        "(Artigos, I/II/III, a)/b), itens 1.1, capítulos e anexos); não substitua "
        "uma convenção por outra.\n"
        "- Para todos os tipos, preserve cabeçalho, ementa/súmula, preâmbulo, "
        "comandos, textos legais fixos, local/data, assinatura, rodapé, tabelas, "
        "anexos, caixa alta, negrito, itálico, fonte, tamanho, alinhamento, "
        "recuos, espaçamentos e quebras de página quando observados.\n"
        "- Para quadros/formulários, use kind \"table\" com table_headers, "
        "table_rows e table_column_widths; preserve a ordem e use marcadores "
        "{{campo}} nas células variáveis.\n"
        "- Linhas de tabela SEM marcador são descartadas na importação (são "
        "dados do exemplo, de outro processo, e não podem ser republicadas): "
        "toda linha que você mantiver precisa usar {{campo}} nas células "
        "variáveis.\n"
        "- Modele a FORMA, não transcreva todas as ocorrências: em tabelas ou "
        "listas repetidas, mantenha cabeçalhos e no máximo duas linhas "
        "representativas com marcadores. Não replique dados repetidos do "
        "documento de referência.\n"
        "- Concordância de gênero: o documento de referência fala de UMA pessoa "
        "(\"a servidora\", \"ao servidor\", \"nomeado\"). Não congele essa "
        "flexão no "
        "texto: declare um campo select \"genero\" com as opções "
        "[\"masculino\", \"feminino\"] e escreva os trechos flexionados como "
        "{{genero:masculino|feminino}} — ex.: "
        "\"Conceder férias {{genero:ao servidor|à servidora}} {{nome_servidor}}\", "
        "\"{{genero:nomeado|nomeada}}\". Vale para artigos, preposições, "
        "pronomes, particípios e para a súmula.\n"
        "- \"document_title\" é o título do ato como ele é publicado, com os "
        "marcadores da própria numeração quando ela existir (ex.: "
        "\"PORTARIA Nº {{numero_portaria}}/{{ano_portaria}}\"); não devolva "
        "apenas o nome genérico do tipo de ato.\n"
        "- \"summary\" é a súmula/ementa do ato. Quando houver um campo de "
        "súmula, o valor deve ser apenas o marcador desse campo (ex.: "
        "\"{{sumula_texto}}\"), nunca o texto de um documento específico. A "
        "súmula descreve objetivamente o que o ato faz e NÃO repete a "
        "fundamentação legal já citada no preâmbulo.\n"
        "- NÃO invente dados; apenas modele a estrutura observada.\n"
        "- Responda apenas o JSON."
    )


def build_user_message(document_type: str, documents: list[tuple[str, str]]) -> str:
    blocks = []
    for idx, (name, text) in enumerate(documents, start=1):
        excerpt = _representative_excerpt(text)
        blocks.append(f"--- DOCUMENTO {idx}: {name} ---\n{excerpt}")
    body = "\n\n".join(blocks)[:MAX_TOTAL_CHARS]
    return (
        f"Tipo de documento esperado: {document_type}.\n"
        "Alguns PDFs são edições completas e podem conter vários tipos de ato. "
        "Use somente os trechos observados, identifique todas as estruturas "
        "relevantes para o tipo esperado, preservando a ordem e a forma de "
        "numeração dos dispositivos, e proponha o modelo documental "
        "estruturado (JSON):\n\n" + body
    )


def _rows_without_sample_data(rows: list[list[str]]) -> list[list[str]]:
    """Remove das tabelas do modelo as linhas de dados do documento-exemplo.

    Uma linha sem nenhum marcador é dado concreto do ato que serviu de
    referência — itens, fornecedores, valores de OUTRO processo. Como o modelo
    insere seus textos fixos em toda minuta, essas linhas seriam publicadas no
    Diário como se fossem do novo ato: informação oficial falsa sobre terceiros.
    O molde da tabela (cabeçalhos, colunas) é preservado e as linhas variáveis —
    as que trazem ``{{campo}}`` — continuam valendo.
    """
    return [row for row in rows if any(extract_markers(str(cell)) for cell in row)]


def _repair_config(data: dict, document_type: str) -> dict:
    """Normaliza a resposta da IA para caber em ``DocumentModelConfig``.

    - Valida escopo, kinds e tipos de campo.
    - Garante que todo marcador tenha um campo declarado (cria campos text).
    - Descarta seções/kinds inválidos em vez de falhar.
    """
    if not isinstance(data, dict):
        raise ValueError("A IA não retornou um objeto JSON.")

    scope = str(data.get("scope_document_type") or document_type).strip().lower()
    if scope not in SCOPE_DOCUMENT_TYPES:
        scope = document_type if document_type in SCOPE_DOCUMENT_TYPES else "outro"

    fields_raw = data.get("fields") if isinstance(data.get("fields"), list) else []
    fields: list[dict] = []
    seen_keys: set[str] = set()
    for f in fields_raw:
        if not isinstance(f, dict):
            continue
        key = str(f.get("key") or "").strip().lower()
        if not _MARKER_KEY.match(key) or key in seen_keys:
            continue
        ftype = str(f.get("type") or "text").strip().lower()
        if ftype not in {t.value for t in FieldType}:
            ftype = "text"
        field = {
            "key": key,
            "label": str(f.get("label") or key),
            "type": ftype,
            "required": bool(f.get("required", False)),
            "options": [str(o) for o in (f.get("options") or []) if str(o).strip()],
        }
        if field["type"] == "select" and not field["options"]:
            field["type"] = "text"
        fields.append(field)
        seen_keys.add(key)

    valid_kinds = {k.value for k in SectionKind}
    # Sub-blocos que não podem ser raiz (o validador rejeitaria a config).
    non_root_kinds = {"paragraph_item", "inciso", "alinea"}
    sections_raw = data.get("sections") if isinstance(data.get("sections"), list) else []
    sections: list[dict] = []
    used_ids: set[str] = set()

    def norm_section(s: dict) -> dict | None:
        if not isinstance(s, dict):
            return None
        kind = str(s.get("kind") or "paragraph").strip().lower()
        if kind not in valid_kinds:
            return None
        sid = str(s.get("id") or "").strip().lower()
        if not _MARKER_KEY.match(sid) or sid in used_ids:
            sid = f"s{uuid.uuid4().hex[:8]}"
        used_ids.add(sid)
        out: dict = {
            "id": sid,
            "kind": kind,
            "text": str(s.get("text") or ""),
            "alignment": str(
                s.get("alignment")
                or ("justify" if kind in {"paragraph", "article"} else "center")
            ),
            "level": int(s.get("level") or 1),
        }
        if s.get("number") is not None:
            out["number"] = str(s["number"])
        for flag in ("fixed_text", "locked", "ai_generated"):
            if s.get(flag):
                out[flag] = True
        out["font_weight"] = "bold" if str(s.get("font_weight") or "normal").lower() == "bold" else "normal"
        out["italic"] = bool(s.get("italic", False))
        out["uppercase"] = bool(s.get("uppercase", False))
        for style_key in ("indent_em", "space_before_mm", "space_after_mm"):
            try:
                out[style_key] = max(0.0, min(float(s.get(style_key, 0) or 0), 30.0))
            except (TypeError, ValueError):
                out[style_key] = 0.0
        if s.get("when_field"):
            out["when_field"] = str(s["when_field"])
            out["when_value"] = str(s.get("when_value") or "")
        if kind == "signature_block":
            entries = s.get("entries") if isinstance(s.get("entries"), list) else []
            clean = [
                {"name": str(e.get("name") or ""), "role": str(e.get("role") or "")}
                for e in entries
                if isinstance(e, dict)
            ]
            out["entries"] = clean or [
                {"name": "{{autoridade_nome}}", "role": "{{autoridade_cargo}}"}
            ]
        if kind == "table":
            headers = s.get("table_headers") if isinstance(s.get("table_headers"), list) else []
            rows = s.get("table_rows") if isinstance(s.get("table_rows"), list) else []
            clean_rows = [
                [str(cell) for cell in row]
                for row in rows
                if isinstance(row, list)
            ]
            column_count = len(headers) or max((len(row) for row in clean_rows), default=0)
            if column_count < 1 or any(len(row) != column_count for row in clean_rows):
                return None
            out["table_headers"] = [str(header) for header in headers]
            out["table_rows"] = _rows_without_sample_data(clean_rows)
            widths = s.get("table_column_widths") if isinstance(s.get("table_column_widths"), list) else []
            if len(widths) == column_count:
                try:
                    parsed_widths = [float(width) for width in widths]
                except (TypeError, ValueError):
                    parsed_widths = []
                if parsed_widths and all(width > 0 for width in parsed_widths):
                    out["table_column_widths"] = parsed_widths
        children_raw = s.get("children") if isinstance(s.get("children"), list) else []
        if kind == "article" and children_raw:
            children = [norm_section(c) for c in children_raw]
            out["children"] = [c for c in children if c]
        return out

    for s in sections_raw:
        normalized = norm_section(s) if isinstance(s, dict) else None
        if normalized and normalized["kind"] not in non_root_kinds:
            sections.append(normalized)

    # Garante campos para todos os marcadores usados.
    texts: list[str] = [str(data.get("document_title") or ""), str(data.get("summary") or "")]
    for s in sections:
        texts.append(str(s.get("text") or ""))
        for e in s.get("entries", []) or []:
            texts.extend([str(e.get("name") or ""), str(e.get("role") or "")])
        texts.extend(str(header) for header in s.get("table_headers", []) or [])
        for row in s.get("table_rows", []) or []:
            texts.extend(str(cell) for cell in row)
    for value in texts:
        for marker in extract_markers(value):
            if marker not in seen_keys:
                fields.append(
                    {
                        "key": marker,
                        "label": marker.replace("_", " ").title(),
                        "type": "text",
                        "required": False,
                        "options": [],
                    }
                )
                seen_keys.add(marker)

    default_purpose = "Modelo aprendido pela IA"
    purpose = str(data.get("purpose") or default_purpose).strip() or default_purpose
    return {
        "purpose": purpose[:200],
        "description": str(data.get("description") or "")[:2000],
        "scope_document_type": scope,
        "document_title": str(data.get("document_title") or "")[:300],
        "summary": str(data.get("summary") or "")[:1000],
        "fields": fields,
        "sections": sections,
    }


async def analyze_documents(
    document_type: str,
    documents: list[tuple[str, str]],
    api_key: str,
    *,
    transport=None,
) -> tuple[DocumentModelConfig, dict]:
    """Analisa os documentos e devolve uma config validada (proposta)."""
    client = DeepSeekClient(api_key, transport=transport)
    messages = [
        {"role": "system", "content": build_system_message()},
        {"role": "user", "content": build_user_message(document_type, documents)},
    ]
    # A tarefa é extração determinística, não redação. Sem raciocínio estendido
    # o modelo devolve a proposta JSON de modo previsível, sem consumir o
    # orçamento que seria necessário ao próprio esquema.
    data, meta = await client.complete_json(
        messages,
        max_tokens=16384,
        disable_thinking=True,
    )
    repaired = _repair_config(data, document_type)
    try:
        config = DocumentModelConfig.model_validate(repaired)
    except ValidationError as exc:
        raise AiInvalidResponseError(
            "A proposta da IA não pôde ser validada como modelo documental."
        ) from exc
    return config, meta


__all__ = [
    "LEARN_PROMPT_VERSION",
    "extract_text",
    "build_system_message",
    "build_user_message",
    "analyze_documents",
]
