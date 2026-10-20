"""Conversão fiel de .docx para HTML do editor (na ordem do documento).

Por que existe: colar do Word perde o que está dentro de **caixas de texto /
formas de desenho** — o navegador recebe só um ``<img src="file:///C:/...">``
apontando para um arquivo temporário do computador do usuário, que ele não
consegue ler (vira imagem quebrada). Avisos de licitação costumam trazer o
quadro de vencedores numa caixa assim.

Lendo o próprio .docx, o conteúdo da caixa (parágrafos e tabelas) entra no
lugar onde ela está, como texto editável. Também preserva o que o importador
antigo perdia: a ordem entre parágrafos e tabelas, o alinhamento e as células
mescladas.

Figuras de verdade (brasão, fotos) não são importadas: o sanitizador não
aceita imagem embutida (``data:``). Elas são contadas e o usuário é avisado.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from html import escape

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
V = "{urn:schemas-microsoft-com:vml}"
WPS = "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}"
WPG = "{http://schemas.microsoft.com/office/word/2010/wordprocessingGroup}"

_ALIGN = {"center": "center", "right": "right", "end": "right", "both": "justify", "distribute": "justify"}


@dataclass
class DocxConversion:
    html: str
    text_boxes: int = 0
    tables: int = 0
    images_skipped: int = 0
    warnings: list[str] = field(default_factory=list)


class _Converter:
    def __init__(self) -> None:
        self.text_boxes = 0
        self.tables = 0
        self.images = 0

    # ── blocos ───────────────────────────────────────────────────────────
    def blocks(self, container) -> list[str]:
        out: list[str] = []
        for child in container:
            tag = child.tag
            if tag == f"{W}p":
                out.extend(self.paragraph(child))
            elif tag == f"{W}tbl":
                out.append(self.table(child))
            elif tag == f"{W}sdt":
                content = child.find(f"{W}sdtContent")
                if content is not None:
                    out.extend(self.blocks(content))
        return out

    def paragraph(self, p) -> list[str]:
        """O parágrafo e, logo depois, o conteúdo das caixas ancoradas nele."""
        inline, boxes = self.inline(p)
        out: list[str] = []
        if inline.replace("<br/>", "").strip():
            style = ""
            jc = p.find(f"{W}pPr/{W}jc")
            align = _ALIGN.get(jc.get(f"{W}val", "")) if jc is not None else None
            if align:
                style = f' style="text-align: {align};"'
            out.append(f"<p{style}>{inline}</p>")
        out.extend(boxes)
        return out

    def inline(self, element) -> tuple[str, list[str]]:
        parts: list[str] = []
        boxes: list[str] = []
        for node in element:
            tag = node.tag
            if tag == f"{W}r":
                text, run_boxes = self.run(node)
                parts.append(text)
                boxes.extend(run_boxes)
            elif tag in (f"{W}hyperlink", f"{W}smartTag", f"{W}ins", f"{W}fldSimple"):
                text, nested = self.inline(node)
                parts.append(text)
                boxes.extend(nested)
            elif tag == f"{W}sdt":
                content = node.find(f"{W}sdtContent")
                if content is not None:
                    text, nested = self.inline(content)
                    parts.append(text)
                    boxes.extend(nested)
        return "".join(parts), boxes

    def run(self, r) -> tuple[str, list[str]]:
        pieces: list[str] = []
        boxes: list[str] = []
        for node in r:
            tag = node.tag
            if tag == f"{W}t":
                pieces.append(escape(node.text or ""))
            elif tag == f"{W}tab":
                pieces.append(" ")
            elif tag in (f"{W}br", f"{W}cr"):
                pieces.append("<br/>")
            elif tag in (f"{W}drawing", f"{W}pict", f"{MC}AlternateContent", f"{W}object"):
                boxes.extend(self.graphic(node))
        text = "".join(pieces)
        if text.strip():
            rpr = r.find(f"{W}rPr")
            if rpr is not None:
                if _on(rpr.find(f"{W}u"), off_values=("none",)):
                    text = f"<u>{text}</u>"
                if _on(rpr.find(f"{W}i")):
                    text = f"<em>{text}</em>"
                if _on(rpr.find(f"{W}b")):
                    text = f"<strong>{text}</strong>"
        return text, boxes

    def graphic(self, node) -> list[str]:
        """Caixas de texto/formas: o conteúdo vira blocos; figuras são contadas.

        ``mc:AlternateContent`` traz a mesma forma duas vezes (Choice moderna e
        Fallback VML) — só a primeira escolha é lida, senão o quadro duplicaria.
        """
        if node.tag == f"{MC}AlternateContent":
            choice = node.find(f"{MC}Choice")
            node = choice if choice is not None else node.find(f"{MC}Fallback")
            if node is None:
                return []
        group = next(node.iter(f"{WPG}wgp"), None)
        if group is not None:
            rebuilt = self.drawn_layout(group)
            if rebuilt is not None:
                return rebuilt
        out: list[str] = []
        boxes = [
            box for box in node.iter(f"{W}txbxContent")
            if not _inside(box, node, f"{W}txbxContent")  # aninhadas: tratadas pela externa
        ]
        for box in boxes:
            self.text_boxes += 1
            out.extend(self.blocks(box))
        if not boxes and (
            any(True for _ in node.iter(f"{A}blip")) or any(True for _ in node.iter(f"{V}imagedata"))
        ):
            self.images += 1
        return out

    # ── quadro desenhado (PDF convertido em Word) ─────────────────────────
    def drawn_layout(self, group) -> list[str] | None:
        """Remonta tabelas desenhadas com caixas de texto + linhas.

        Documentos convertidos de PDF (comum em avisos e homologações gerados
        por sistemas de compras) não têm tabela: cada linha de texto é uma
        caixa posicionada e as bordas são linhas soltas. As linhas definem a
        grade; cada caixa vai para a célula onde está; o que fica fora de
        qualquer célula vira parágrafo, na ordem vertical.
        Devolve ``None`` quando o grupo não tem grade (segue o caminho comum).
        """
        texts, horizontals, verticals = [], [], []
        for sp in group.findall(f"{WPS}wsp"):
            off = sp.find(f"{WPS}spPr/{A}xfrm/{A}off")
            ext = sp.find(f"{WPS}spPr/{A}xfrm/{A}ext")
            if off is None or ext is None:
                continue
            x, y = int(off.get("x", 0)), int(off.get("y", 0))
            cx, cy = int(ext.get("cx", 0)), int(ext.get("cy", 0))
            box = sp.find(f"{WPS}txbx/{W}txbxContent")
            if box is not None:
                html = " ".join(t for t in (self.inline(p)[0].strip() for p in box.iter(f"{W}p")) if t)
                if html:
                    texts.append({"x": x, "y": y, "cx": cx, "cy": cy, "html": html})
            elif cx > 5 * max(cy, 1):
                horizontals.append((y, x, x + cx))
            elif cy > 5 * max(cx, 1):
                verticals.append((x, y, y + cy))
        if not texts or len(horizontals) < 2 or len(verticals) < 2:
            return None

        tol = max(2000, min(t["cy"] for t in texts) // 4)

        def same(a, b):
            return abs(a - b) <= tol

        ys = _cluster(sorted(h[0] for h in horizontals), tol)
        # Faixas entre linhas horizontais consecutivas; em cada uma, as
        # verticais que a atravessam inteira são as divisas das células.
        bands = []
        for top, bottom in zip(ys, ys[1:]):
            xs = _cluster(sorted(
                v[0] for v in verticals if v[1] <= top + tol and v[2] >= bottom - tol
            ), tol)
            if len(xs) >= 2:
                bands.append({"top": top, "bottom": bottom, "xs": xs})
        if not bands:
            return None

        # Faixas coladas viram a mesma tabela.
        tables: list[list[dict]] = []
        for band in bands:
            if tables and same(tables[-1][-1]["bottom"], band["top"]):
                tables[-1].append(band)
            else:
                tables.append([band])

        placed: set[int] = set()
        items: list[tuple[int, str]] = []
        for bands_of_table in tables:
            union = _cluster(sorted(x for b in bands_of_table for x in b["xs"]), tol)
            rows_html = []
            rows_cells = []
            for band in bands_of_table:
                cells = []
                for left, right in zip(band["xs"], band["xs"][1:]):
                    inside = [
                        i for i, t in enumerate(texts)
                        if i not in placed
                        and left - tol <= t["x"] < right - tol
                        and band["top"] - tol <= t["y"] + t["cy"] / 2 <= band["bottom"] + tol
                    ]
                    placed.update(inside)
                    span = sum(1 for a in union if left - tol <= a < right - tol)
                    cells.append((span, _join_lines([texts[i] for i in inside], tol)))
                rows_cells.append(cells[0] if len(cells) == 1 else (0, ""))
                rows_html.append("<tr>" + "".join(
                    f'<td{f" colspan=\"{span}\"" if span > 1 else ""} style="border: 1px solid #000;">{html}</td>'
                    for span, html in cells
                ) + "</tr>")
            if len(union) <= 2:
                # Uma coluna só = parágrafo emoldurado, não tabela: vira
                # parágrafos (como tabela, a diagramação do ato a desmanchava
                # e a juntava ao fecho/assinatura).
                for band, (_, html) in zip(bands_of_table, (row for row in rows_cells)):
                    if html:
                        items.append((band["top"], f"<p>{html}</p>"))
                continue
            self.tables += 1
            items.append((
                bands_of_table[0]["top"],
                '<table style="border-collapse: collapse; width: 100%;">' + "".join(rows_html) + "</table>",
            ))

        loose = [t for i, t in enumerate(texts) if i not in placed]
        for line in _group_lines(loose, tol):
            items.append((line[0]["y"], f"<p>{_join_lines(line, tol)}</p>"))
        self.text_boxes += 1
        return [html for _, html in sorted(items, key=lambda item: item[0])]

    # ── tabelas ──────────────────────────────────────────────────────────
    def table(self, tbl) -> str:
        self.tables += 1
        rows = tbl.findall(f"{W}tr")
        grid: list[list[dict]] = []
        for tr in rows:
            cells = []
            for tc in tr.findall(f"{W}tc"):
                tcpr = tc.find(f"{W}tcPr")
                span = 1
                vmerge = None
                if tcpr is not None:
                    gs = tcpr.find(f"{W}gridSpan")
                    if gs is not None:
                        span = int(gs.get(f"{W}val", "1") or 1)
                    vm = tcpr.find(f"{W}vMerge")
                    if vm is not None:
                        vmerge = vm.get(f"{W}val", "continue")
                cells.append({"tc": tc, "span": span, "vmerge": vmerge, "rowspan": 1})
            grid.append(cells)

        # Mesclagem vertical: a célula "restart" absorve as "continue" abaixo,
        # casando pela coluna de grade (soma dos colspans à esquerda).
        def columns(row):
            col = 0
            for cell in row:
                yield col, cell
                col += cell["span"]

        for r_index, row in enumerate(grid):
            for col, cell in columns(row):
                if cell["vmerge"] != "restart":
                    continue
                for below in grid[r_index + 1:]:
                    match = next((c for cc, c in columns(below) if cc == col), None)
                    if match is None or match["vmerge"] != "continue":
                        break
                    cell["rowspan"] += 1

        html = ['<table style="border-collapse: collapse; width: 100%;">']
        for row in grid:
            html.append("<tr>")
            for cell in row:
                if cell["vmerge"] == "continue":
                    continue
                attrs = ""
                if cell["span"] > 1:
                    attrs += f' colspan="{cell["span"]}"'
                if cell["rowspan"] > 1:
                    attrs += f' rowspan="{cell["rowspan"]}"'
                html.append(f'<td{attrs} style="border: 1px solid #000;">{self.cell(cell["tc"])}</td>')
            html.append("</tr>")
        html.append("</table>")
        return "".join(html)

    def cell(self, tc) -> str:
        lines: list[str] = []
        for child in tc:
            if child.tag == f"{W}p":
                text, boxes = self.inline(child)
                if text.strip():
                    lines.append(text)
                lines.extend(_strip_block_tags(b) for b in boxes)
            elif child.tag == f"{W}tbl":
                # Tabela dentro de célula: rara; o texto é mantido linha a linha.
                for p in child.iter(f"{W}p"):
                    text, _ = self.inline(p)
                    if text.strip():
                        lines.append(text)
        return "<br/>".join(lines)


def _cluster(values: list[int], tol: int) -> list[int]:
    """Valores ordenados → um representante por grupo de valores próximos."""
    out: list[int] = []
    for value in values:
        if not out or value - out[-1] > tol:
            out.append(value)
    return out


def _group_lines(texts: list[dict], tol: int) -> list[list[dict]]:
    """Caixas na mesma altura formam uma linha (ordenadas da esquerda)."""
    lines: list[list[dict]] = []
    for t in sorted(texts, key=lambda t: (t["y"], t["x"])):
        if lines and abs(lines[-1][0]["y"] - t["y"]) <= tol:
            lines[-1].append(t)
        else:
            lines.append([t])
    return [sorted(line, key=lambda t: t["x"]) for line in lines]


def _join_lines(texts: list[dict], tol: int) -> str:
    """Junta as linhas de uma célula num texto corrido (o editor refaz a
    quebra). Na mesma linha, caixas encostadas ("1" + ". SOM") colam sem
    espaço."""
    parts: list[str] = []
    for line in _group_lines(texts, tol):
        chunk = ""
        prev = None
        for t in line:
            glued = prev is not None and t["x"] <= prev["x"] + prev["cx"]
            chunk += ("" if not chunk or glued else " ") + t["html"]
            prev = t
        parts.append(chunk)
    return " ".join(parts)


def _on(element, off_values: tuple[str, ...] = ()) -> bool:
    if element is None:
        return False
    val = (element.get(f"{W}val") or "").lower()
    return val not in ("0", "false", "off", *off_values)


def _inside(element, stop, tag) -> bool:
    parent = element.getparent()
    while parent is not None and parent is not stop:
        if parent.tag == tag:
            return True
        parent = parent.getparent()
    return False


def _strip_block_tags(html: str) -> str:
    import re

    return re.sub(r"</?(?:p|table|tr|td)\b[^>]*>", " ", html).strip()


def docx_to_html(content: bytes) -> DocxConversion:
    from docx import Document

    from app.core.html_sanitizer import sanitize_html

    document = Document(io.BytesIO(content))
    converter = _Converter()
    html = "\n".join(converter.blocks(document.element.body))
    result = DocxConversion(
        html=sanitize_html(html),
        text_boxes=converter.text_boxes,
        tables=converter.tables,
        images_skipped=converter.images,
    )
    if converter.images:
        result.warnings.append(
            f"{converter.images} figura(s) do Word (brasão, foto, assinatura escaneada) "
            "não foram importadas — o texto e as tabelas, sim."
        )
    return result
