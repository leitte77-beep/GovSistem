"""Word → HTML do Novo ato: caixa de texto com o quadro de vencedores.

Caso real (homologação do Pregão 51/2026): o quadro estava numa caixa de
desenho; colado no editor virava imagem quebrada (file:///C:/...).
"""

import io

from docx import Document
from lxml import etree

from app.services.docx_html import docx_to_html

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NSDECL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)

BOX_CONTENT = """
<w:txbxContent>
  <w:p><w:r><w:t>Fornecedor e Itens declarados Vencedores (cfe. Cotação):</w:t></w:r></w:p>
  <w:tbl>
    <w:tr>
      <w:tc><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>lote</w:t></w:r></w:p></w:tc>
      <w:tc><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Descrição</w:t></w:r></w:p></w:tc>
      <w:tc><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Valor_Total</w:t></w:r></w:p></w:tc>
    </w:tr>
    <w:tr>
      <w:tc><w:p><w:r><w:t>1</w:t></w:r></w:p></w:tc>
      <w:tc><w:p><w:r><w:t>SOM (P.A, MONITOR E BACKLINE)</w:t></w:r></w:p></w:tc>
      <w:tc><w:p><w:r><w:t>48.524,9700</w:t></w:r></w:p></w:tc>
    </w:tr>
  </w:tbl>
  <w:tbl>
    <w:tr><w:tc><w:tcPr><w:gridSpan w:val="3"/></w:tcPr><w:p><w:r><w:t>Fornecedor: F&amp;V SHOWS E ESTRUTURAS LTDA</w:t></w:r></w:p></w:tc></w:tr>
  </w:tbl>
</w:txbxContent>
"""


def _docx_with_text_box() -> bytes:
    doc = Document()
    title = doc.add_paragraph()
    title.add_run("TERMO DE ADJUDICAÇÃO E HOMOLOGAÇÃO").bold = True
    title.paragraph_format.alignment = 1  # center
    anchor = doc.add_paragraph("Adjudicando e Homologando o procedimento, à empresa:")
    # Mesma forma em Choice (wps) e Fallback (VML), como o Word grava.
    shape = etree.fromstring(
        f"""<w:r {NSDECL}><mc:AlternateContent>
          <mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData>
            <wps:wsp><wps:txbx>{BOX_CONTENT}</wps:txbx></wps:wsp>
          </a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>
          <mc:Fallback><w:pict><v:shape><v:textbox>{BOX_CONTENT}</v:textbox></v:shape></w:pict></mc:Fallback>
        </mc:AlternateContent></w:r>"""
    )
    anchor._p.append(shape)
    closing = doc.add_paragraph("Município de Farol, 05/10/2026.")
    closing.paragraph_format.alignment = 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_text_box_content_becomes_editable_in_place_without_duplicates():
    result = docx_to_html(_docx_with_text_box())
    html = result.html

    assert result.text_boxes == 1
    assert result.tables == 2
    assert "<img" not in html
    # Conteúdo da caixa como texto/tabela, uma vez só (Choice, não Fallback).
    assert html.count("SOM (P.A, MONITOR E BACKLINE)") == 1
    assert html.count("Fornecedor e Itens declarados Vencedores") == 1
    # Na ordem do documento: âncora → caixa → fecho.
    order = [html.index(t) for t in ("Adjudicando", "Fornecedor e Itens", "48.524,9700", "Município de Farol")]
    assert order == sorted(order)
    # Formatação e alinhamento preservados.
    assert "<strong>lote</strong>" in html
    assert 'colspan="3"' in html
    assert "text-align: center" in html.split("Município de Farol")[0].rsplit("<p", 1)[1]


def test_vertical_merge_and_pictures_are_reported():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Lote 1"
    table.cell(0, 1).text = "a"
    table.cell(1, 1).text = "b"
    buf = io.BytesIO()
    doc.save(buf)
    result = docx_to_html(buf.getvalue())
    assert 'rowspan="2"' in result.html
    assert result.html.count("Lote 1") == 1
    assert result.images_skipped == 0


def test_real_pdf_converted_word_with_drawn_table_is_rebuilt():
    """Arquivo real (homologação do Pregão 51/2026, convertido de PDF): não há
    tabela — 26 caixas de texto (uma por linha de texto) + 80 linhas de borda.
    O quadro precisa voltar como tabela, célula por célula."""
    from pathlib import Path

    data = (Path(__file__).parent / "fixtures" / "homologacao_pregao_51_quadro_desenhado.docx").read_bytes()
    result = docx_to_html(data)
    html = result.html

    # O quadro principal é tabela; a caixa de uma coluna ("concat/Fornecedor")
    # vira parágrafos.
    assert result.tables == 1
    assert html.count("<table") == 1
    first = html.split("<table", 2)[1]
    rows = first.split("<tr")[1:]
    assert len(rows) == 2
    assert rows[0].count("<td") == 9 and rows[1].count("<td") == 9
    assert "1. SOM (P.A, MONITOR E BACKLINE) , conforme tr" in rows[1]
    assert "F&amp;V SHOWS E ESTRUTURAS LTDA" in rows[1]
    assert "48.524,9700" in rows[1]
    # Texto fora da grade continua parágrafo e na ordem do documento.
    order = [html.index(t) for t in ("Adjudicando", "Fornecedor e Itens", "lote", "concat", "TERMO DE CONVÊNIO", "Município de Farol")]
    assert order == sorted(order)
    assert html.count("Valor Total Homologado") == 1  # Choice, não o Fallback VML
    assert "<p>Fornecedor: F&amp;V SHOWS" in html


def test_real_file_closing_keeps_only_place_and_date():
    """O fecho engolia "TERMO DE CONVÊNIO Nº 0525/2026", o "Objeto…" e o
    "Fornecedor… CNPJ …/0001-07" (qualquer número de 4 dígitos contava como
    data)."""
    from pathlib import Path

    from app.semantic.parser import parse_document

    data = (Path(__file__).parent / "fixtures" / "homologacao_pregao_51_quadro_desenhado.docx").read_bytes()
    doc = parse_document(html=docx_to_html(data).html, title="PROCESSOS DE COMPRA", document_type="licitacao")
    signature = [b for b in doc.blocks if b.type == "signature_block"][0]
    assert signature.entries[0].location == "Município de Farol, 05/10/2026."
    body = " ".join(getattr(b, "content", "") or "" for b in doc.blocks)
    assert "TERMO DE CONVÊNIO" in body and "Valor Total Homologado" in body
