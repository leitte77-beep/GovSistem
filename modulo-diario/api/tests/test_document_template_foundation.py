"""Testes da Fase 1 — fundações do construtor visual de modelos documentais.

Cobre: layout visual versionado (modelo visual), identidade institucional,
herança (parent_model_id), soft delete, biblioteca de blocos reutilizáveis e
autorização/isolamento por organização.
"""

from __future__ import annotations

import io
import json
import uuid

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient, MockTransport, Response
from pydantic import ValidationError

from app.core.auth import get_current_user
from app.core.database import get_db
from app.document_model.body_html import blocks_to_html
from app.document_model.layout import DocumentLayout, default_layout
from app.document_model.learning import analyze_documents, build_system_message
from app.main import app
from app.models.organization import Organization
from app.models.user import User
from app.semantic.schemas import SemanticDocument


class _RoleName:
    def __init__(self, name: str):
        self.name = name


class _RoleObj:
    def __init__(self, name: str):
        self.role = _RoleName(name)


class _UserProxy:
    def __init__(self, user: User, roles: list[str]):
        self._user = user
        self._roles = [_RoleObj(r) for r in roles]

    @property
    def id(self):
        return self._user.id

    @property
    def organization_id(self):
        return self._user.organization_id

    @property
    def email(self):
        return self._user.email

    @property
    def user_roles(self):
        return self._roles


@pytest_asyncio.fixture
async def ctx(db_session):
    org = Organization(name="Prefeitura Teste", slug="pf-f1", cnpj="12345678000194")
    db_session.add(org)
    await db_session.flush()

    rows = {}
    for label, roles in {"admin": ["ADMIN"], "autor": ["AUTOR"], "consulta": ["CONSULTA"]}.items():
        u = User(
            name=label,
            email=f"{label}_{uuid.uuid4().hex[:8]}@test",
            organization_id=org.id,
            is_active=True,
        )
        db_session.add(u)
        rows[label] = u
    await db_session.flush()

    class C:
        org_id = org.id
        admin = _UserProxy(rows["admin"], ["ADMIN"])
        autor = _UserProxy(rows["autor"], ["AUTOR"])
        consulta = _UserProxy(rows["consulta"], ["CONSULTA"])

    yield C()
    app.dependency_overrides.clear()


@pytest_asyncio.fixture
async def api_client(ctx, db_session):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        ac_ = ac

        def bind(user):
            app.dependency_overrides[get_db] = lambda: db_session
            app.dependency_overrides[get_current_user] = lambda: user
            return ac_

        yield bind


def _config_dict(**overrides) -> dict:
    cfg = {
        "purpose": "Concessão de férias",
        "scope_document_type": "portaria",
        "document_title": "PORTARIA DE FÉRIAS",
        "summary": "Concede férias ao servidor {{servidor}}.",
        "fields": [{"key": "servidor", "label": "Servidor", "type": "text", "required": True}],
        "sections": [
            {"id": "art1", "kind": "article", "text": "Concede férias a {{servidor}}."}
        ],
    }
    cfg.update(overrides)
    return cfg


# ── Layout visual (modelo visual) ───────────────────────────────────────────


def test_default_layout_is_valid():
    layout = default_layout()
    assert layout["page_size"] == "A4"
    assert layout["orientation"] == "portrait"
    assert layout["margins"]["top"] == 20


def test_layout_rejects_invalid_values():
    with pytest.raises(ValidationError):
        DocumentLayout(page_size="B5")
    with pytest.raises(ValidationError):
        DocumentLayout(margins={"top": 999})
    with pytest.raises(ValidationError):
        DocumentLayout(body_font={"color": "azul"})


# ── Layout versionado + herança ─────────────────────────────────────────────


async def test_create_model_with_layout_and_parent(api_client, ctx):
    client = api_client(ctx.admin)
    parent = await client.post(
        "/api/v1/document-models",
        json={"name": "Portaria Base", "slug": "portaria-base", "config": _config_dict()},
    )
    assert parent.status_code == 201
    parent_id = parent.json()["id"]

    layout = {"page_size": "A4", "body_font": {"family": "Arial", "size": 11}}
    child = await client.post(
        "/api/v1/document-models",
        json={
            "name": "Portaria de Férias",
            "slug": "portaria-ferias-f1",
            "config": _config_dict(),
            "layout": layout,
            "parent_model_id": parent_id,
        },
    )
    assert child.status_code == 201
    assert child.json()["parent_model_id"] == parent_id

    detail = await client.get(f"/api/v1/document-models/{child.json()['id']}/versions/1")
    assert detail.status_code == 200
    assert detail.json()["layout"]["body_font"]["family"] == "Arial"


async def test_create_model_rejects_invalid_layout(api_client, ctx):
    client = api_client(ctx.admin)
    resp = await client.post(
        "/api/v1/document-models",
        json={
            "name": "X",
            "slug": "x-layout",
            "config": _config_dict(),
            "layout": {"page_size": "B5"},
        },
    )
    assert resp.status_code == 422


async def test_create_model_rejects_foreign_parent(api_client, ctx):
    client = api_client(ctx.admin)
    resp = await client.post(
        "/api/v1/document-models",
        json={
            "name": "X",
            "slug": "x-parent",
            "config": _config_dict(),
            "parent_model_id": str(uuid.uuid4()),
        },
    )
    assert resp.status_code == 404


async def test_new_version_keeps_layout(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "M", "slug": "m-ver", "config": _config_dict()},
    )
    model_id = created.json()["id"]
    resp = await client.post(
        f"/api/v1/document-models/{model_id}/versions",
        json={"config": _config_dict(), "layout": {"orientation": "landscape"}},
    )
    assert resp.status_code == 201
    assert resp.json()["layout"]["orientation"] == "landscape"


# ── Autosave: edição de rascunho ────────────────────────────────────────────


async def test_patch_draft_version_updates_config_and_layout(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Autosave", "slug": "autosave-f2", "config": _config_dict()},
    )
    model_id = created.json()["id"]
    resp = await client.patch(
        f"/api/v1/document-models/{model_id}/versions/1",
        json={
            "config": _config_dict(purpose="Nova finalidade"),
            "layout": {"body_font": {"family": "Arial"}},
            "change_reason": "autosave",
        },
    )
    assert resp.status_code == 200
    assert resp.json()["layout"]["body_font"]["family"] == "Arial"
    assert resp.json()["change_reason"] == "autosave"

    detail = await client.get(f"/api/v1/document-models/{model_id}")
    assert detail.json()["purpose"] == "Nova finalidade"


async def test_patch_active_version_is_rejected(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Ativo", "slug": "ativo-f2", "config": _config_dict()},
    )
    model_id = created.json()["id"]
    await client.post(f"/api/v1/document-models/{model_id}/versions/1/submit")
    await client.post(f"/api/v1/document-models/{model_id}/versions/1/approve")
    resp = await client.patch(
        f"/api/v1/document-models/{model_id}/versions/1",
        json={"config": _config_dict(purpose="Alterado")},
    )
    assert resp.status_code == 409


# ── Render HTML/PDF unificado (Fase 3) ──────────────────────────────────────


def test_body_html_renders_article_children_from_canonical_content():
    document = SemanticDocument(
        title="Estrutura legal",
        blocks=[
            {
                "type": "article",
                "caput": "Caput",
                "paragraphs": [{"content": "Texto do parágrafo"}],
                "incisos": [{"number": "I", "content": "Texto do inciso"}],
                "alineas": [{"number": "a", "content": "Texto da alínea"}],
            }
        ],
    )

    rendered = blocks_to_html(document)

    assert "Texto do parágrafo" in rendered
    assert "Texto do inciso" in rendered
    assert "Texto da alínea" in rendered


async def test_render_html_and_pdf(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Render", "slug": "render-f3", "config": _config_dict()},
    )
    model_id = created.json()["id"]

    html = await client.post(
        f"/api/v1/document-models/{model_id}/versions/1/render",
        json={"values": {"servidor": "João da Silva"}},
    )
    assert html.status_code == 200
    body = html.json()
    assert body["complete"] is True
    assert "João da Silva" in body["html"]
    assert "@page" in body["html"]

    pdf = await client.post(
        f"/api/v1/document-models/{model_id}/versions/1/render-pdf",
        json={"values": {"servidor": "João da Silva"}},
    )
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")


async def test_render_rejects_unknown_field(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Render2", "slug": "render2-f3", "config": _config_dict()},
    )
    model_id = created.json()["id"]
    resp = await client.post(
        f"/api/v1/document-models/{model_id}/versions/1/render",
        json={"values": {"inexistente": "x"}},
    )
    assert resp.status_code == 422


async def test_section_flags_roundtrip(api_client, ctx):
    client = api_client(ctx.admin)
    cfg = _config_dict()
    cfg["sections"][0]["fixed_text"] = True
    cfg["sections"][0]["locked"] = True
    cfg["sections"][0]["ai_generated"] = False
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Flags", "slug": "flags-f3", "config": cfg},
    )
    assert created.status_code == 201
    model_id = created.json()["id"]
    detail = await client.get(f"/api/v1/document-models/{model_id}/versions/1")
    assert detail.json()["config"]["sections"][0]["locked"] is True
    assert detail.json()["config"]["sections"][0]["fixed_text"] is True


# ── Assinatura: autoridade/certificado/posição (item 21) ────────────────────


def _signature_config(position: str = "right") -> dict:
    return _config_dict(
        sections=[
            {
                "id": "sig1",
                "kind": "signature_block",
                "alignment": "center",
                "entries": [
                    {
                        "name": "Oclécio de Freitas Meneses",
                        "role": "Prefeito Municipal",
                        "authority_id": str(uuid.uuid4()),
                        "credential_id": str(uuid.uuid4()),
                        "position": position,
                    }
                ],
            }
        ]
    )


async def test_signature_entry_authority_credential_position(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Sig", "slug": "sig-f1", "config": _signature_config("right")},
    )
    assert created.status_code == 201, created.text
    detail = await client.get(
        f"/api/v1/document-models/{created.json()['id']}/versions/1"
    )
    entry = detail.json()["config"]["sections"][0]["entries"][0]
    assert entry["position"] == "right"
    assert entry["authority_id"]
    assert entry["credential_id"]


async def test_signature_entry_rejects_invalid_position(api_client, ctx):
    client = api_client(ctx.admin)
    resp = await client.post(
        "/api/v1/document-models",
        json={"name": "Sig2", "slug": "sig2-f1", "config": _signature_config("diagonal")},
    )
    assert resp.status_code == 422


async def test_signature_render_applies_entry_position(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Sig3", "slug": "sig3-f1", "config": _signature_config("left")},
    )
    model_id = created.json()["id"]
    html = await client.post(
        f"/api/v1/document-models/{model_id}/versions/1/render", json={"values": {}}
    )
    assert html.status_code == 200
    assert "text-align:left" in html.json()["html"]
    assert "OCLÉCIO DE FREITAS MENESES" in html.json()["html"]


# ── Aprender com documentos (Fase 4) ────────────────────────────────────────

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_bytes(text: str) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _create_model(client, slug: str) -> str:
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Treino", "slug": slug, "config": _config_dict()},
    )
    assert created.status_code == 201
    return created.json()["id"]


async def test_training_file_upload_list_toggle_delete(api_client, ctx):
    client = api_client(ctx.admin)
    model_id = await _create_model(client, "treino-f4")
    up = await client.post(
        f"/api/v1/document-models/{model_id}/training-files",
        files={"file": ("portaria.docx", _docx_bytes("Conceder férias ao servidor."), DOCX_MIME)},
    )
    assert up.status_code == 201
    file_id = up.json()["id"]
    assert up.json()["used_by_ai"] is False
    assert up.json()["status"] == "analyzed"

    listed = await client.get(f"/api/v1/document-models/{model_id}/training-files")
    assert any(f["id"] == file_id for f in listed.json())

    on = await client.patch(
        f"/api/v1/document-models/{model_id}/training-files/{file_id}",
        json={"used_by_ai": True},
    )
    assert on.status_code == 200 and on.json()["used_by_ai"] is True

    off = await client.patch(
        f"/api/v1/document-models/{model_id}/training-files/{file_id}",
        json={"used_by_ai": False},
    )
    assert off.json()["used_by_ai"] is False

    deleted = await client.delete(
        f"/api/v1/document-models/{model_id}/training-files/{file_id}"
    )
    assert deleted.status_code == 204
    listed2 = await client.get(f"/api/v1/document-models/{model_id}/training-files")
    assert all(f["id"] != file_id for f in listed2.json())


async def test_training_file_rejects_non_pdf_docx(api_client, ctx):
    client = api_client(ctx.admin)
    model_id = await _create_model(client, "treino-bad")
    resp = await client.post(
        f"/api/v1/document-models/{model_id}/training-files",
        files={"file": ("dados.csv", b"a,b\n1,2\n", "text/csv")},
    )
    assert resp.status_code == 422


async def test_propose_requires_marked_documents(api_client, ctx):
    client = api_client(ctx.admin)
    model_id = await _create_model(client, "treino-empty")
    resp = await client.post(f"/api/v1/document-models/{model_id}/training-files/propose")
    assert resp.status_code == 409


async def test_propose_without_ai_key_returns_typed_failure(api_client, ctx):
    client = api_client(ctx.admin)
    model_id = await _create_model(client, "treino-nokey")
    up = await client.post(
        f"/api/v1/document-models/{model_id}/training-files",
        files={"file": ("p.docx", _docx_bytes("RESOLVE: conceder férias."), DOCX_MIME)},
    )
    file_id = up.json()["id"]
    await client.patch(
        f"/api/v1/document-models/{model_id}/training-files/{file_id}",
        json={"used_by_ai": True},
    )
    resp = await client.post(f"/api/v1/document-models/{model_id}/training-files/propose")
    assert resp.status_code == 200
    body = resp.json()
    assert body["ok"] is False
    assert body["status"] == "not_configured"


def test_analyze_documents_repairs_missing_field_markers():
    proposal = {
        "purpose": "Concessão de férias",
        "scope_document_type": "portaria",
        "document_title": "PORTARIA",
        "summary": "Concede férias a {{servidor}}.",
        "fields": [],
        "sections": [
            {"id": "art1", "kind": "article", "text": "Conceder {{dias}} dias a {{servidor}}."}
        ],
    }
    transport = MockTransport(
        lambda req: Response(
            200,
            json={
                "choices": [{"message": {"content": json.dumps(proposal)}}],
                "usage": {"total_tokens": 10},
                "model": "deepseek-v4-flash",
            },
        )
    )
    import asyncio

    config, _ = asyncio.run(
        analyze_documents("portaria", [("p.docx", "texto")], "sk-test", transport=transport)
    )
    keys = {f.key for f in config.fields}
    assert {"servidor", "dias"} <= keys


def test_analyze_documents_drops_root_subblocks():
    proposal = {
        "purpose": "Concessão",
        "scope_document_type": "portaria",
        "document_title": "PORTARIA",
        "summary": "",
        "fields": [],
        "sections": [
            {"id": "i1", "kind": "inciso", "text": "I - inválido na raiz"},
            {"id": "a1", "kind": "article", "text": "Art. 1 Concede."},
        ],
    }
    transport = MockTransport(
        lambda req: Response(
            200,
            json={"choices": [{"message": {"content": json.dumps(proposal)}}]},
        )
    )
    import asyncio

    config, _ = asyncio.run(
        analyze_documents("portaria", [("p.docx", "texto")], "sk-test", transport=transport)
    )
    assert [s.kind.value for s in config.sections] == ["article"]


def test_learning_prompt_limits_repeated_table_rows():
    prompt = build_system_message()
    assert "no máximo duas linhas representativas" in prompt


# ── Soft delete ─────────────────────────────────────────────────────────────

async def test_soft_delete_hides_model(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-models",
        json={"name": "Apagar", "slug": "apagar-f1", "config": _config_dict()},
    )
    model_id = created.json()["id"]
    assert (await client.delete(f"/api/v1/document-models/{model_id}")).status_code == 204
    assert (await client.get(f"/api/v1/document-models/{model_id}")).status_code == 404
    listing = await client.get("/api/v1/document-models")
    assert all(m["id"] != model_id for m in listing.json())


# ── Identidade institucional ────────────────────────────────────────────────


async def test_institution_get_and_patch(api_client, ctx):
    client = api_client(ctx.admin)
    got = await client.get("/api/v1/settings/institution")
    assert got.status_code == 200
    assert got.json()["name"] == "Prefeitura Teste"

    patched = await client.patch(
        "/api/v1/settings/institution",
        json={
            "state": "CE",
            "address_city": "Farol",
            "phone": "(88) 3663-0000",
            "site": "https://farol.ce.gov.br",
            "institutional_layout": {"page_size": "A4"},
        },
    )
    assert patched.status_code == 200
    body = patched.json()
    assert body["state"] == "CE"
    assert body["address_city"] == "Farol"
    assert body["institutional_layout"]["page_size"] == "A4"


async def test_institution_patch_requires_permission(api_client, ctx):
    client = api_client(ctx.consulta)
    resp = await client.patch("/api/v1/settings/institution", json={"state": "CE"})
    assert resp.status_code == 403


# ── Biblioteca de blocos ────────────────────────────────────────────────────


async def test_blocks_crud(api_client, ctx):
    client = api_client(ctx.admin)
    created = await client.post(
        "/api/v1/document-model-blocks",
        json={
            "name": "Registre-se e Publique-se",
            "kind": "command",
            "content_json": {"text": "x"},
        },
    )
    assert created.status_code == 201
    block_id = created.json()["id"]

    listed = await client.get("/api/v1/document-model-blocks")
    assert any(b["id"] == block_id for b in listed.json())

    patched = await client.patch(
        f"/api/v1/document-model-blocks/{block_id}", json={"description": "rodapé padrão"}
    )
    assert patched.status_code == 200
    assert patched.json()["description"] == "rodapé padrão"

    assert (await client.delete(f"/api/v1/document-model-blocks/{block_id}")).status_code == 204
    listed2 = await client.get("/api/v1/document-model-blocks")
    assert all(b["id"] != block_id for b in listed2.json())


async def test_block_duplicate_name_conflicts(api_client, ctx):
    client = api_client(ctx.admin)
    payload = {"name": "Paço Municipal", "kind": "text", "content_json": {}}
    assert (await client.post("/api/v1/document-model-blocks", json=payload)).status_code == 201
    assert (await client.post("/api/v1/document-model-blocks", json=payload)).status_code == 409


async def test_consultation_cannot_create_block(api_client, ctx):
    client = api_client(ctx.consulta)
    resp = await client.post(
        "/api/v1/document-model-blocks",
        json={"name": "X", "kind": "text", "content_json": {}},
    )
    assert resp.status_code == 403
