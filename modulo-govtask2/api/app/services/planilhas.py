"""Planilhas Excel (xlsx) dos pedidos.

Duas saídas: o relatório de um pedido, com uma aba por assunto, e a lista
consolidada, com o mesmo recorte dos filtros da tela. Tudo é leitura — a
planilha é uma fotografia, não uma porta de escrita.
"""

import io
import re
import uuid
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.pedido import PedidoDetalhe, PedidoLista


def _texto_simples(texto):
    """Tira a marcação do editor (negrito, itálico, listas) — o Excel é texto puro."""
    if not texto:
        return texto
    linhas = []
    for linha in texto.splitlines():
        linha = re.sub(r"^\s*(?:[-*]|\d+[.)])\s+", "", linha)
        linha = re.sub(r"\*\*(.+?)\*\*", r"\1", linha)
        linha = re.sub(r"(?<![\w])_([^_\n]+)_(?![\w])", r"\1", linha)
        linhas.append(linha)
    return "\n".join(linhas)


def _normalizar(valor):
    """Tipos que o Excel não engole: datetime com fuso, Decimal, UUID, Enum."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, uuid.UUID):
        return str(valor)
    value = getattr(valor, "value", None)
    if value is not None and not isinstance(valor, (str, int, float, bool)):
        return value
    return valor

CABECALHO = PatternFill("solid", fgColor="1F3A5F")
BRANCO = Font(bold=True, color="FFFFFF")
LARGURA_MAX = 60


def _aba(ws, cabecalhos: list[str], linhas: list[list]):
    linhas = [[_normalizar(v) for v in linha] for linha in linhas]
    ws.append(cabecalhos)
    for celula in ws[1]:
        celula.fill = CABECALHO
        celula.font = BRANCO
        celula.alignment = Alignment(vertical="center")
    for linha in linhas:
        ws.append(linha)

    for indice, titulo in enumerate(cabecalhos, start=1):
        comprimento = max(
            [len(str(titulo))]
            + [len(str(linha[indice - 1])) for linha in linhas if linha[indice - 1] is not None]
        )
        ws.column_dimensions[get_column_letter(indice)].width = min(
            LARGURA_MAX, max(12, comprimento + 2)
        )
    for linha in ws.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def _texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor)


def _bool_pt(valor: bool) -> str:
    return "Sim" if valor else "Não"


def workbook_do_pedido(pedido: PedidoDetalhe) -> io.BytesIO:
    wb = Workbook()

    dados = wb.active
    dados.title = "Dados"
    _aba(
        dados,
        ["Campo", "Valor"],
        [
            ["Número", pedido.numero],
            ["Título", pedido.titulo],
            ["Tipo", pedido.tipo],
            ["Situação", pedido.situacao],
            ["Prioridade", pedido.prioridade],
            ["Origem", pedido.origem],
            ["Quem conseguiu", pedido.origem_nome],
            ["Tarefa atual", pedido.tarefa_atual],
            ["Setor atual", pedido.setor_atual],
            ["Prazo atual", pedido.prazo_atual],
            ["Dias de atraso", pedido.dias_de_atraso],
            ["Valor previsto", pedido.valor_previsto],
            ["Valor liberado", pedido.valor_liberado],
            ["Valor pago", pedido.valor_pago],
            ["Protocolo (número)", pedido.protocolo_externo],
            ["Protocolo (sistema)", pedido.protocolo_sistema],
            ["Protocolo (órgão)", pedido.protocolo_orgao],
            ["Protocolo (data)", pedido.protocolo_data],
            ["Aberto por", pedido.criado_por.name if pedido.criado_por else None],
            ["Aberto em", pedido.created_at],
            ["Concluído em", pedido.concluido_em],
            ["Motivo do cancelamento", pedido.motivo_cancelamento],
            ["Descrição", _texto_simples(pedido.descricao)],
        ],
    )

    tramitacao = wb.create_sheet("Tramitação")
    _aba(
        tramitacao,
        [
            "Ordem", "Assunto", "Setor", "Status", "Responsável", "Mencionados",
            "Prazo", "Assumido em", "Devolvido em", "Resultado",
            "Complemento pedido", "Complemento respondido", "Transferências",
        ],
        [
            [
                e.ordem,
                e.assunto,
                e.setor,
                e.status,
                e.responsavel.name if e.responsavel else "",
                ", ".join(p.name for p in e.participantes),
                e.prazo,
                e.assumido_em,
                e.devolvido_em,
                _texto_simples(e.resultado),
                e.complemento_pedido,
                e.complemento_resposta,
                e.transferencias,
            ]
            for e in sorted(pedido.encaminhamentos, key=lambda x: x.ordem)
        ],
    )

    anexos = wb.create_sheet("Anexos")
    _aba(
        anexos,
        ["Origem", "Categoria", "Nome", "Descrição", "Versão", "Enviado por", "Tamanho (bytes)", "Enviado em"],
        [
            [
                f"Encaminhamento #{a.encaminhamento_id}"
                if a.encaminhamento_id
                else ("Medição" if a.medicao_id else "Pedido"),
                a.categoria,
                a.nome_original,
                a.descricao,
                a.versao,
                a.enviado_por.name if a.enviado_por else "",
                a.tamanho_bytes,
                a.created_at,
            ]
            for a in pedido.anexos
        ],
    )

    if pedido.tipo == "OBRA":
        medicoes = wb.create_sheet("Medições")
        _aba(
            medicoes,
            ["Número", "Período início", "Período fim", "Valor", "% executado", "Responsável", "Observação", "Fotos"],
            [
                [
                    m.numero,
                    m.periodo_inicio,
                    m.periodo_fim,
                    m.valor,
                    m.percentual_executado,
                    m.responsavel.name if m.responsavel else "",
                    m.observacao,
                    ", ".join(f.nome_original for f in m.fotos),
                ]
                for m in pedido.medicoes
            ],
        )

    historico = wb.create_sheet("Histórico")
    _aba(
        historico,
        ["Data", "Tipo", "Autor", "Texto"],
        [
            [a.created_at, a.tipo, a.autor_nome, a.texto]
            for a in sorted(pedido.andamentos, key=lambda x: x.created_at)
        ],
    )

    return _salvar(wb)


def workbook_da_lista(itens: list[PedidoLista]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pedidos"
    _aba(
        ws,
        [
            "Número", "Título", "Tipo", "Situação", "Prioridade", "Origem",
            "Quem conseguiu", "Setor atual", "Responsável", "Tarefa atual",
            "Prazo", "Dias de atraso", "Valor previsto", "Saúde", "Próxima ação",
        ],
        [
            [
                p.numero,
                p.titulo,
                p.tipo,
                p.situacao,
                p.prioridade,
                p.origem,
                p.origem_nome,
                p.setor_atual,
                p.responsavel_atual.name if p.responsavel_atual else "",
                p.tarefa_atual,
                p.prazo_atual,
                p.dias_de_atraso,
                p.valor_previsto,
                p.saude,
                p.proxima_acao,
            ]
            for p in itens
        ],
    )
    return _salvar(wb)


def _salvar(wb: Workbook) -> io.BytesIO:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
