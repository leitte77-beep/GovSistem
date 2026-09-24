"""Relatórios: o pedido e a lista em Excel, e a página de PDF (impressão)."""

import io

import pytest
from openpyxl import load_workbook

CARRO = {"titulo": "Carro do deputado", "tipo": "AQUISICAO"}
OBRA = {"titulo": "Pavimentação do Bairro Novo", "tipo": "OBRA"}


async def _abrir(cliente, dados=None):
    resposta = await cliente.post("/pedidos", json=dados or CARRO)
    assert resposta.status_code == 201, resposta.text
    return resposta.json()


async def _encaminhar(cliente, pedido_id):
    resposta = await cliente.post(
        f"/pedidos/{pedido_id}/encaminhar",
        json={"setor": "ENGENHARIA", "assunto": "Projeto e orçamento"},
    )
    assert resposta.status_code == 200, resposta.text
    return resposta.json()


def _workbook(conteudo: bytes):
    return load_workbook(io.BytesIO(conteudo))


@pytest.mark.asyncio
async def test_relatorio_do_pedido_em_excel(como, assessor):
    async with como(assessor) as cliente:
        pedido = await _abrir(cliente)
        await _encaminhar(cliente, pedido["id"])
        resposta = await cliente.get(f"/pedidos/{pedido['id']}/relatorio.xlsx")

    assert resposta.status_code == 200
    assert "spreadsheetml" in resposta.headers["content-type"]
    assert resposta.content[:2] == b"PK"
    wb = _workbook(resposta.content)
    assert {"Dados", "Tramitação", "Anexos", "Histórico"} <= set(wb.sheetnames)


@pytest.mark.asyncio
async def test_relatorio_de_obra_traz_aba_de_medicoes(como, assessor, engenheiro):
    async with como(assessor) as cliente:
        pedido = await _abrir(cliente, OBRA)
        corpo = await _encaminhar(cliente, pedido["id"])
    enc_id = corpo["encaminhamento_atual"]["id"]

    async with como(engenheiro) as cliente:
        await cliente.post(f"/pedidos/{pedido['id']}/encaminhamentos/{enc_id}/assumir")
        await cliente.post(
            f"/pedidos/{pedido['id']}/encaminhamentos/{enc_id}/medicoes",
            json={"valor": "1000.00", "percentual_executado": "10.00"},
        )

    async with como(assessor) as cliente:
        resposta = await cliente.get(f"/pedidos/{pedido['id']}/relatorio.xlsx")

    wb = _workbook(resposta.content)
    assert "Medições" in wb.sheetnames
    aba = wb["Medições"]
    assert aba.max_row == 2  # cabeçalho + uma medição


@pytest.mark.asyncio
async def test_descricao_formatada_sai_texto_puro_no_excel(como, assessor):
    async with como(assessor) as cliente:
        pedido = await _abrir(
            cliente,
            {**CARRO, "descricao": "**Urgente**\n- item _um_\n1. item dois"},
        )
        resposta = await cliente.get(f"/pedidos/{pedido['id']}/relatorio.xlsx")

    aba = _workbook(resposta.content)["Dados"]
    valores = {
        aba.cell(row=r, column=1).value: aba.cell(row=r, column=2).value
        for r in range(1, aba.max_row + 1)
    }
    assert valores["Descrição"] == "Urgente\nitem um\nitem dois"


@pytest.mark.asyncio
async def test_lista_consolidada_em_excel_respeita_filtros(como, assessor):
    async with como(assessor) as cliente:
        await _abrir(cliente, {**CARRO, "titulo": "Pedido de aquisição"})
        await _abrir(cliente, {**CARRO, "titulo": "Pedido de outro tipo", "tipo": "OUTRO"})
        resposta = await cliente.get("/pedidos/exportar.xlsx?tipo=AQUISICAO")

    assert resposta.status_code == 200
    wb = _workbook(resposta.content)
    aba = wb["Pedidos"]
    # Cabeçalho + só a aquisição.
    assert aba.max_row == 2
